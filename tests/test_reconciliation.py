"""
Milestone 19 proof-point tests: reconciliation, rollforward, and supporting schedule infrastructure.

Tests
-----
 1. Create reconciliation with correct variance calculation
 2. Reconciliation status transitions (not_started → in_progress → prepared → reviewed)
 3. Rollforward carries prior period ending balance as new opening balance
 4. Tie-out marks in_tolerance when variance <= tolerance
 5. Tie-out marks out_of_tolerance when variance exceeds tolerance
 6. Reconciliation lines accumulate correctly
 7. Support reference links external system reference to reconciliation
 8. Reviewer separation: preparer cannot also review
 9. Cash rollforward formula: opening + inflows - outflows
10. RE rollforward formula: beginning RE + net income - dividends
11. Draft preview balance can be updated on reconciliation
12. Reconciliation export workbook is labeled correctly
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models.account import Account
from app.models.accounting_period import AccountingPeriod
from app.models.entity import Entity
import datetime

from app.services.organization_service import create_organization
from app.services.reconciliation_service import (
    ReconciliationStateError,
    ReviewerSeparationError,
    add_reconciliation_line,
    add_support_reference,
    build_cash_rollforward,
    build_re_rollforward,
    create_reconciliation,
    get_reconciliation,
    get_reconciliation_lines,
    get_support_references,
    rollforward_reconciliation,
    transition_status,
    update_reconciliation_balances,
)
from app.services.export_service import build_reconciliation_workbook, workbook_to_bytes


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def db():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


@pytest.fixture(scope="module")
def fixtures(db):
    org = create_organization(db, name="Recon Corp", slug="recon-corp")
    db.flush()

    ent = Entity(
        organization_id=org.id, code="ENT1", name="Recon Entity",
        entity_type="operating", currency="USD",
    )
    db.add(ent)
    db.flush()

    period1 = AccountingPeriod(
        entity_id=ent.id, period_name="Jan 2024",
        start_date=datetime.date(2024, 1, 1), end_date=datetime.date(2024, 1, 31),
        fiscal_year=2024, fiscal_period=1, period_type="monthly",
    )
    period2 = AccountingPeriod(
        entity_id=ent.id, period_name="Feb 2024",
        start_date=datetime.date(2024, 2, 1), end_date=datetime.date(2024, 2, 29),
        fiscal_year=2024, fiscal_period=2, period_type="monthly",
    )
    db.add_all([period1, period2])
    db.flush()

    cash_acct = Account(
        entity_id=ent.id, account_number="1000", account_name="Cash",
        account_type="asset", normal_balance="debit",
    )
    re_acct = Account(
        entity_id=ent.id, account_number="3900", account_name="Retained Earnings",
        account_type="equity", normal_balance="credit",
    )
    db.add_all([cash_acct, re_acct])
    db.flush()

    db.commit()
    return {
        "org": org, "ent": ent,
        "period1": period1, "period2": period2,
        "cash_acct": cash_acct, "re_acct": re_acct,
    }


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_1_create_reconciliation_variance(db, fixtures):
    """Create reconciliation with correct variance calculation."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]
    period1 = fixtures["period1"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        period_id=period1.id,
        reconciliation_type="bank",
        official_balance=Decimal("10000.00"),
        supporting_balance=Decimal("9850.00"),
        tolerance_amount=Decimal("200.00"),
    )
    db.commit()

    assert recon.id is not None
    assert recon.variance_amount == Decimal("150.00")
    assert recon.tie_out_status == "in_tolerance"
    assert recon.reconciliation_type == "bank"


def test_2_status_transitions(db, fixtures):
    """Status transitions from not_started through reviewed."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("5000"),
        supporting_balance=Decimal("5000"),
    )
    db.commit()
    assert recon.status == "not_started"

    transition_status(db, recon.id, "in_progress")
    db.commit()
    assert recon.status == "in_progress"

    transition_status(db, recon.id, "prepared", user_id=1)
    db.commit()
    assert recon.status == "prepared"
    assert recon.preparer_user_id == 1
    assert recon.prepared_at is not None

    transition_status(db, recon.id, "reviewed", user_id=2, comment="Looks good")
    db.commit()
    assert recon.status == "reviewed"
    assert recon.reviewer_user_id == 2
    assert recon.reviewer_comment == "Looks good"


def test_3_rollforward_carries_opening_balance(db, fixtures):
    """Rollforward carries prior period closing balance as new opening balance."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]
    period1 = fixtures["period1"]
    period2 = fixtures["period2"]

    prior = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        period_id=period1.id,
        official_balance=Decimal("12000"),
        supporting_balance=Decimal("12000"),
    )
    db.commit()
    transition_status(db, prior.id, "in_progress")
    transition_status(db, prior.id, "prepared", user_id=3)
    transition_status(db, prior.id, "reviewed", user_id=4)
    db.commit()

    new_recon = rollforward_reconciliation(
        db, prior.id, new_period_id=period2.id, new_official_balance=Decimal("13500")
    )
    db.commit()

    assert new_recon.rollforward_opening_balance == Decimal("12000")
    assert new_recon.rollforward_closing_balance == Decimal("12000")  # opening + 0 adjustments
    assert new_recon.official_balance == Decimal("13500")
    assert prior.status == "rolled_forward"


def test_4_tie_out_in_tolerance(db, fixtures):
    """Tie-out marks in_tolerance when |variance| <= tolerance."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("1000"),
        supporting_balance=Decimal("999.50"),
        tolerance_amount=Decimal("1.00"),
    )
    db.commit()

    assert recon.variance_amount == Decimal("0.50")
    assert recon.tie_out_status == "in_tolerance"


def test_5_tie_out_out_of_tolerance(db, fixtures):
    """Tie-out marks out_of_tolerance when |variance| > tolerance."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("1000"),
        supporting_balance=Decimal("950"),
        tolerance_amount=Decimal("10"),
    )
    db.commit()

    assert recon.variance_amount == Decimal("50")
    assert recon.tie_out_status == "out_of_tolerance"


def test_6_reconciliation_lines_accumulate(db, fixtures):
    """Reconciliation lines accumulate in line_number order."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("3000"),
        supporting_balance=Decimal("3000"),
    )
    db.commit()

    line1 = add_reconciliation_line(
        db, recon.id,
        description="Opening bank balance",
        source_type="bank",
        debit=Decimal("2500"),
        credit=Decimal("0"),
    )
    line2 = add_reconciliation_line(
        db, recon.id,
        description="Deposit in transit",
        source_type="gl",
        debit=Decimal("500"),
        credit=Decimal("0"),
        is_reconciling_item=True,
        reconciling_notes="Deposit cleared Feb 2",
    )
    db.commit()

    lines = get_reconciliation_lines(db, recon.id)
    assert len(lines) >= 2
    assert line1.line_number < line2.line_number
    assert line1.balance == Decimal("2500")
    assert line2.is_reconciling_item is True


def test_7_support_reference_links(db, fixtures):
    """Support reference links external system reference to reconciliation."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
    )
    db.commit()

    ref = add_support_reference(
        db, recon.id,
        reference_type="external_system",
        external_ref="BANK-STMT-2024-01",
        description="January 2024 bank statement",
        added_by_user_id=1,
    )
    db.commit()

    refs = get_support_references(db, recon.id)
    assert len(refs) >= 1
    assert any(r.external_ref == "BANK-STMT-2024-01" for r in refs)
    assert ref.reference_type == "external_system"
    assert ref.added_at is not None


def test_8_reviewer_separation(db, fixtures):
    """Preparer cannot also be the reviewer."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
    )
    db.commit()
    transition_status(db, recon.id, "in_progress")
    transition_status(db, recon.id, "prepared", user_id=99)
    db.commit()

    with pytest.raises(ReviewerSeparationError):
        transition_status(db, recon.id, "reviewed", user_id=99)


def test_9_cash_rollforward_formula(db, fixtures):
    """Cash rollforward: opening + inflows - outflows = closing."""
    lines = build_cash_rollforward(
        opening_balance=Decimal("50000"),
        inflows=Decimal("15000"),
        outflows=Decimal("8000"),
    )
    closing_line = next(l for l in lines if l.is_subtotal)
    assert closing_line.amount == Decimal("57000")


def test_10_re_rollforward_formula(db, fixtures):
    """RE rollforward: beginning RE + net income - dividends = ending RE."""
    lines = build_re_rollforward(
        beginning_re=Decimal("100000"),
        net_income=Decimal("25000"),
        dividends=Decimal("5000"),
    )
    closing_line = next(l for l in lines if l.is_subtotal)
    assert closing_line.amount == Decimal("120000")


def test_11_draft_preview_balance_update(db, fixtures):
    """Draft preview balance can be updated on a reconciliation."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("20000"),
        supporting_balance=Decimal("20000"),
    )
    db.commit()

    updated = update_reconciliation_balances(
        db, recon.id,
        draft_preview_balance=Decimal("21500"),
        variance_explanation="Draft adjustment for audit topside",
    )
    db.commit()

    assert updated.draft_preview_balance == Decimal("21500")
    assert updated.variance_explanation == "Draft adjustment for audit topside"


def test_12_reconciliation_export_labeled(db, fixtures):
    """Reconciliation export workbook has correct label in cover sheet."""
    org = fixtures["org"]
    ent = fixtures["ent"]
    cash = fixtures["cash_acct"]

    recon = create_reconciliation(
        db,
        organization_id=org.id,
        entity_id=ent.id,
        account_id=cash.id,
        official_balance=Decimal("7500"),
        supporting_balance=Decimal("7500"),
        notes="Monthly bank recon",
    )
    db.commit()

    add_reconciliation_line(
        db, recon.id,
        description="GL ending balance",
        source_type="gl",
        debit=Decimal("7500"),
    )
    ref = add_support_reference(
        db, recon.id,
        reference_type="workpaper",
        external_ref="WP-001",
        description="Bank statement",
    )
    db.commit()

    lines = get_reconciliation_lines(db, recon.id)
    support = get_support_references(db, recon.id)

    wb = build_reconciliation_workbook(recon, lines, support)
    data = workbook_to_bytes(wb)

    wb2 = openpyxl.load_workbook(io.BytesIO(data))
    sheet_names = wb2.sheetnames
    assert "Reconciliation" in sheet_names
    assert "Reconciliation Lines" in sheet_names
    assert "Support References" in sheet_names

    cover = wb2["Reconciliation"]
    # Check that cover sheet has the RECONCILIATION label
    cell_values = [cover.cell(row=r, column=1).value for r in range(1, 12)]
    assert any("RECONCILIATION" in str(v) for v in cell_values if v)
