"""
Taxonomy library API — list/get/clone taxonomies, edit nodes, export, and
manage account-to-taxonomy mappings (multi-taxonomy aware).

System taxonomies are immutable; users must clone before editing.
"""
from __future__ import annotations

import csv
import io
import json

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.api.deps import get_db
from app.api.schemas import (
    AccountTaxonomyMappingBulkRequest,
    AccountTaxonomyMappingCreate,
    AccountTaxonomyMappingOut,
    TaxonomyCloneRequest,
    TaxonomyDetailOut,
    TaxonomyNodeCreate,
    TaxonomyNodeOut,
    TaxonomyNodeTreeOut,
    TaxonomyNodeUpdate,
    TaxonomyOut,
)
from app.models.taxonomy import (
    AccountTaxonomyMapping,
    Taxonomy,
    TaxonomyNode,
)
from app.services import taxonomy_library_service as svc
from app.services.taxonomy_library_service import (
    TaxonomyImmutableError,
    TaxonomyNotFoundError,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    _HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    _HAS_OPENPYXL = False


router = APIRouter(tags=["taxonomies"])


EXPORT_COLUMNS = [
    "taxonomy_name",
    "code",
    "parent_code",
    "name",
    "description",
    "statement_type",
    "financial_statement_section",
    "normal_balance",
    "sort_order",
    "level",
    "gaap_reference",
    "ifrs_reference",
    "xbrl_tag",
    "cash_flow_classification",
    "consolidation_treatment",
    "kpi_eligible",
    "industry",
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _load_taxonomy_or_404(db: Session, taxonomy_id: int) -> Taxonomy:
    try:
        return svc.get_taxonomy(db, taxonomy_id)
    except TaxonomyNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


def _export_rows(db: Session, taxonomy: Taxonomy) -> list[dict]:
    """Flatten taxonomy nodes into export dicts (one per node)."""
    nodes = (
        db.query(TaxonomyNode)
        .filter_by(taxonomy_id=taxonomy.id)
        .order_by(TaxonomyNode.level, TaxonomyNode.sort_order)
        .all()
    )
    parent_codes = {n.id: n.code for n in nodes}
    rows: list[dict] = []
    for n in nodes:
        rows.append({
            "taxonomy_name": taxonomy.name,
            "code": n.code,
            "parent_code": parent_codes.get(n.parent_id) if n.parent_id else "",
            "name": n.name,
            "description": n.description or "",
            "statement_type": n.statement_type or "",
            "financial_statement_section": n.financial_statement_section or "",
            "normal_balance": n.normal_balance or "",
            "sort_order": n.sort_order,
            "level": n.level,
            "gaap_reference": n.gaap_reference or "",
            "ifrs_reference": n.ifrs_reference or "",
            "xbrl_tag": n.xbrl_tag or "",
            "cash_flow_classification": n.cash_flow_classification or "",
            "consolidation_treatment": n.consolidation_treatment or "",
            "kpi_eligible": n.kpi_eligible,
            "industry": n.industry or "",
        })
    return rows


# ---------------------------------------------------------------------------
# Taxonomy CRUD
# ---------------------------------------------------------------------------

@router.get("", response_model=list[TaxonomyOut])
def list_taxonomies(
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
):
    return svc.list_taxonomies(db, include_inactive=include_inactive)


@router.get("/{taxonomy_id}", response_model=TaxonomyDetailOut)
def get_taxonomy(taxonomy_id: int, db: Session = Depends(get_db)):
    tx = _load_taxonomy_or_404(db, taxonomy_id)
    node_count = db.query(TaxonomyNode).filter_by(taxonomy_id=tx.id).count()
    return TaxonomyDetailOut(
        id=tx.id,
        code=tx.code,
        name=tx.name,
        description=tx.description,
        industry=tx.industry,
        version=tx.version,
        is_system=tx.is_system,
        parent_taxonomy_id=tx.parent_taxonomy_id,
        is_active=tx.is_active,
        node_count=node_count,
    )


@router.get("/{taxonomy_id}/tree", response_model=list[TaxonomyNodeTreeOut])
def get_taxonomy_tree(taxonomy_id: int, db: Session = Depends(get_db)):
    _load_taxonomy_or_404(db, taxonomy_id)
    return svc.get_taxonomy_tree(db, taxonomy_id)


@router.get("/{taxonomy_id}/nodes", response_model=list[TaxonomyNodeOut])
def list_taxonomy_nodes(
    taxonomy_id: int,
    search: str | None = Query(None),
    statement_type: str | None = Query(None),
    is_active: bool | None = Query(None),
    db: Session = Depends(get_db),
):
    _load_taxonomy_or_404(db, taxonomy_id)
    q = db.query(TaxonomyNode).filter_by(taxonomy_id=taxonomy_id)
    if statement_type:
        q = q.filter(TaxonomyNode.statement_type == statement_type)
    if is_active is not None:
        q = q.filter(TaxonomyNode.is_active.is_(is_active))
    if search:
        like = f"%{search.lower()}%"
        from sqlalchemy import func, or_
        q = q.filter(or_(
            func.lower(TaxonomyNode.name).like(like),
            func.lower(TaxonomyNode.code).like(like),
        ))
    return q.order_by(TaxonomyNode.level, TaxonomyNode.sort_order).all()


@router.post("/{taxonomy_id}/clone", response_model=TaxonomyOut, status_code=201)
def clone_taxonomy(
    taxonomy_id: int,
    body: TaxonomyCloneRequest,
    db: Session = Depends(get_db),
):
    _load_taxonomy_or_404(db, taxonomy_id)
    try:
        clone = svc.clone_taxonomy(db, taxonomy_id, body.name, body.code)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return clone


# ---------------------------------------------------------------------------
# Node CRUD (only on editable taxonomies)
# ---------------------------------------------------------------------------

@router.post("/{taxonomy_id}/nodes", response_model=TaxonomyNodeOut, status_code=201)
def create_node(
    taxonomy_id: int,
    body: TaxonomyNodeCreate,
    db: Session = Depends(get_db),
):
    _load_taxonomy_or_404(db, taxonomy_id)
    try:
        node = svc.create_node(db, taxonomy_id, body.model_dump(exclude_unset=False))
    except TaxonomyImmutableError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return node


@router.patch("/nodes/{node_id}", response_model=TaxonomyNodeOut)
def update_node(
    node_id: int,
    body: TaxonomyNodeUpdate,
    db: Session = Depends(get_db),
):
    try:
        node = svc.update_node(db, node_id, body.model_dump(exclude_unset=True))
    except TaxonomyNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    except TaxonomyImmutableError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return node


@router.post("/nodes/{node_id}/deactivate", response_model=TaxonomyNodeOut)
def deactivate_node(node_id: int, db: Session = Depends(get_db)):
    try:
        node = svc.deactivate_node(db, node_id)
    except TaxonomyNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    except TaxonomyImmutableError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return node


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

@router.get("/{taxonomy_id}/export/csv")
def export_csv(taxonomy_id: int, db: Session = Depends(get_db)):
    tx = _load_taxonomy_or_404(db, taxonomy_id)
    rows = _export_rows(db, tx)

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    buf.seek(0)
    filename = f"{tx.code}_taxonomy.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{taxonomy_id}/export/excel")
def export_excel(taxonomy_id: int, db: Session = Depends(get_db)):
    if not _HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    tx = _load_taxonomy_or_404(db, taxonomy_id)
    rows = _export_rows(db, tx)

    wb = Workbook()
    ws = wb.active
    ws.title = (tx.code or "Taxonomy")[:31]
    bold = Font(bold=True)
    for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = bold
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row[col_name])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"{tx.code}_taxonomy.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{taxonomy_id}/export/json")
def export_json(taxonomy_id: int, db: Session = Depends(get_db)):
    tx = _load_taxonomy_or_404(db, taxonomy_id)
    tree = svc.get_taxonomy_tree(db, taxonomy_id)
    payload = {
        "code": tx.code,
        "name": tx.name,
        "description": tx.description,
        "industry": tx.industry,
        "version": tx.version,
        "is_system": tx.is_system,
        "nodes": tree,
    }
    body = json.dumps(payload, indent=2, default=str)
    filename = f"{tx.code}_taxonomy.json"
    return StreamingResponse(
        iter([body]),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Account → Taxonomy node mappings
# ---------------------------------------------------------------------------

@router.post("/mappings", response_model=AccountTaxonomyMappingOut, status_code=201)
def create_mapping(
    body: AccountTaxonomyMappingCreate,
    db: Session = Depends(get_db),
):
    return svc.map_account(
        db,
        account_id=body.account_id,
        taxonomy_id=body.taxonomy_id,
        taxonomy_node_id=body.taxonomy_node_id,
        mapping_source=body.mapping_source,
        mapping_type=body.mapping_type,
        confidence_score=body.confidence_score,
        is_primary=body.is_primary,
        mapped_by=body.mapped_by,
    )


@router.post("/mappings/bulk")
def bulk_create_mappings(
    body: AccountTaxonomyMappingBulkRequest,
    db: Session = Depends(get_db),
):
    count = svc.bulk_map_accounts(
        db,
        [m.model_dump() for m in body.mappings],
    )
    return {"count": count}


@router.get("/mappings/account/{account_id}")
def list_account_mappings(account_id: int, db: Session = Depends(get_db)):
    mappings = (
        db.query(AccountTaxonomyMapping)
        .filter_by(account_id=account_id)
        .all()
    )
    out: list[dict] = []
    for m in mappings:
        tx = db.query(Taxonomy).filter_by(id=m.taxonomy_id).first()
        node = db.query(TaxonomyNode).filter_by(id=m.taxonomy_node_id).first()
        out.append({
            "id": m.id,
            "account_id": m.account_id,
            "taxonomy_id": m.taxonomy_id,
            "taxonomy": {
                "id": tx.id, "code": tx.code, "name": tx.name, "is_system": tx.is_system,
            } if tx else None,
            "taxonomy_node_id": m.taxonomy_node_id,
            "node": {
                "id": node.id, "code": node.code, "name": node.name,
                "statement_type": node.statement_type,
                "financial_statement_section": node.financial_statement_section,
            } if node else None,
            "mapping_type": m.mapping_type,
            "mapping_source": m.mapping_source,
            "confidence_score": m.confidence_score,
            "is_primary": m.is_primary,
        })
    return out


@router.delete("/mappings/{mapping_id}", status_code=204)
def delete_mapping(mapping_id: int, db: Session = Depends(get_db)):
    svc.delete_account_mapping(db, mapping_id)
    return None
