"""Tests for Sprint 3.14 — Quarterly Review Generator service."""

import io
from decimal import Decimal

import pytest

from app.services.quarterly_review_service import (
    _pct_change,
    _fmt_amount,
    _fmt_pct,
    _build_financial_changes,
    export_markdown,
    export_excel,
)


# ---------------------------------------------------------------------------
# _pct_change
# ---------------------------------------------------------------------------

class TestPctChange:
    def test_positive_change(self):
        assert _pct_change("110", "100") == Decimal("10.0")

    def test_negative_change(self):
        assert _pct_change("90", "100") == Decimal("-10.0")

    def test_zero_prior_returns_none(self):
        assert _pct_change("100", "0") is None

    def test_none_current_treated_as_zero(self):
        result = _pct_change(None, "100")
        assert result == Decimal("-100.0")

    def test_rounding(self):
        result = _pct_change("133", "300")
        # (133-300)/300*100 = -55.666... → -55.7
        assert result == Decimal("-55.7")


# ---------------------------------------------------------------------------
# _fmt_amount
# ---------------------------------------------------------------------------

class TestFmtAmount:
    def test_millions(self):
        assert _fmt_amount("2500000") == "$2.50M"

    def test_thousands(self):
        assert _fmt_amount("25000") == "$25.0K"

    def test_small_amount(self):
        assert _fmt_amount("500") == "$500"

    def test_none_returns_dash(self):
        assert _fmt_amount(None) == "—"

    def test_empty_string_returns_dash(self):
        assert _fmt_amount("") == "—"

    def test_negative_millions(self):
        result = _fmt_amount("-1500000")
        assert "M" in result


# ---------------------------------------------------------------------------
# _fmt_pct
# ---------------------------------------------------------------------------

class TestFmtPct:
    def test_positive(self):
        assert _fmt_pct(Decimal("10.0")) == "+10.0%"

    def test_negative(self):
        assert _fmt_pct(Decimal("-5.5")) == "-5.5%"

    def test_none_returns_dash(self):
        assert _fmt_pct(None) == "—"


# ---------------------------------------------------------------------------
# _build_financial_changes
# ---------------------------------------------------------------------------

def _make_diag(**overrides) -> dict:
    base = {
        "revenue": "500000",
        "cogs": "300000",
        "gross_profit": "200000",
        "gross_margin_pct": "40.0",
        "total_expenses": "150000",
        "net_income": "50000",
        "cash": "100000",
        "accounts_receivable": "80000",
        "inventory": "60000",
        "total_current_assets": "240000",
        "total_current_liabilities": "120000",
        "total_assets": "500000",
        "total_liabilities": "250000",
        "total_equity": "250000",
        "total_debt": "200000",
        "working_capital": "120000",
        "current_ratio": "2.0",
        "debt_to_equity": "0.8",
        "roa": "10.0",
    }
    base.update(overrides)
    return base


class TestBuildFinancialChanges:
    def test_direction_increase(self):
        cur = _make_diag(revenue="600000")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        rev = next(r for r in rows if r["metric"] == "revenue")
        assert rev["direction"] == "increase"

    def test_direction_decrease(self):
        cur = _make_diag(revenue="400000")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        rev = next(r for r in rows if r["metric"] == "revenue")
        assert rev["direction"] == "decrease"

    def test_significant_above_materiality(self):
        cur = _make_diag(revenue="600000")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        rev = next(r for r in rows if r["metric"] == "revenue")
        assert rev["significant"] is True

    def test_not_significant_below_materiality(self):
        cur = _make_diag(revenue="500500")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        rev = next(r for r in rows if r["metric"] == "revenue")
        assert rev["significant"] is False

    def test_pct_metric_formatting(self):
        cur = _make_diag(gross_margin_pct="42.0")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        gm = next(r for r in rows if r["metric"] == "gross_margin_pct")
        assert gm["current_value"] == "42.0%"
        assert gm["prior_value"] == "40.0%"
        assert "pp" in gm["change_amount"]

    def test_ratio_metric_formatting(self):
        cur = _make_diag(current_ratio="2.50")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        cr = next(r for r in rows if r["metric"] == "current_ratio")
        assert "x" in cr["current_value"]
        assert "x" in cr["prior_value"]

    def test_display_order_respected(self):
        cur = _make_diag()
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        metrics = [r["metric"] for r in rows]
        assert metrics.index("revenue") < metrics.index("net_income")
        assert metrics.index("net_income") < metrics.index("cash")

    def test_change_pct_none_for_pct_metrics(self):
        cur = _make_diag(gross_margin_pct="42.0")
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        gm = next(r for r in rows if r["metric"] == "gross_margin_pct")
        assert gm["change_pct"] is None

    def test_unchanged_direction(self):
        cur = _make_diag()
        pri = _make_diag()
        rows = _build_financial_changes(cur, pri, Decimal("1000"))
        rev = next(r for r in rows if r["metric"] == "revenue")
        assert rev["direction"] == "unchanged"


# ---------------------------------------------------------------------------
# Shared sample report dict
# ---------------------------------------------------------------------------

def _sample_report() -> dict:
    return {
        "metadata": {
            "entity_id": 1,
            "entity_name": "Test Company LLC",
            "current_period": {"id": 1, "name": "Q1 2024", "start": "2024-01-01", "end": "2024-03-31"},
            "comparison_period": {"id": 2, "name": "Q4 2023", "start": "2023-10-01", "end": "2023-12-31"},
            "generated_at": "2024-04-01T00:00:00Z",
            "materiality_threshold": "1000",
        },
        "executive_summary": {
            "overall_risk": "moderate",
            "total_issues": 2,
            "issues_by_severity": {"critical": 0, "high": 1, "moderate": 1, "low": 0, "informational": 0},
            "key_findings": [
                "Revenue increased +20.0% period-over-period.",
                "1 high-severity issue identified.",
            ],
            "narrative": "Test narrative for Q1 2024.",
        },
        "key_financial_changes": [
            {
                "metric": "revenue",
                "label": "Revenue",
                "current_value": "$600.0K",
                "prior_value": "$500.0K",
                "change_amount": "$100.0K",
                "change_pct": "+20.0%",
                "direction": "increase",
                "significant": True,
            },
        ],
        "significant_variances": [
            {
                "metric": "revenue",
                "label": "Revenue",
                "current_value": "$600.0K",
                "prior_value": "$500.0K",
                "change_amount": "$100.0K",
                "change_pct": "+20.0%",
                "direction": "increase",
                "significant": True,
            },
        ],
        "triggered_issues": [
            {
                "issue_code": "REV_001",
                "category": "revenue_recognition",
                "severity": "high",
                "title": "Revenue Growth Exceeds Industry",
                "description": "Revenue grew 20% vs expected 5%.",
                "detection_trigger": "revenue_pct_change=20.0 > threshold=15.0",
                "supporting_metrics": {"revenue_pct_change": "20.0"},
                "management_questions": ["What drove the revenue increase?"],
                "procedure_items": ["Review revenue contracts."],
                "aje_items": [],
                "status": "open",
            },
        ],
        "management_questions": [
            "What drove the revenue increase?",
            "Are there any unusual one-time items?",
        ],
        "suggested_procedures": [
            {
                "issue_code": "REV_001",
                "title": "Revenue Growth Exceeds Industry",
                "severity": "high",
                "items": ["Review revenue contracts.", "Inspect supporting invoices."],
            }
        ],
        "suggested_adjustments": [],
        "advisor_notes": "ENGAGEMENT NOTES — TEST COMPANY LLC — Q1 2024\n\nMateriality: $1,000",
    }


# ---------------------------------------------------------------------------
# export_markdown
# ---------------------------------------------------------------------------

class TestExportMarkdown:
    def test_contains_all_eight_sections(self):
        md = export_markdown(_sample_report())
        for section in [
            "## 1. Executive Summary",
            "## 2. Key Financial Changes",
            "## 3. Significant Variances",
            "## 4. Triggered Accounting Issues",
            "## 5. Suggested Questions for Management",
            "## 6. Suggested Procedures",
            "## 7. Suggested Adjusting Journal Entries",
            "## 8. Advisor Notes",
        ]:
            assert section in md, f"Missing section: {section}"

    def test_header_contains_entity_and_period(self):
        md = export_markdown(_sample_report())
        assert "Test Company LLC" in md
        assert "Q1 2024" in md

    def test_financial_table_row(self):
        md = export_markdown(_sample_report())
        assert "$600.0K" in md
        assert "+20.0%" in md

    def test_management_questions_numbered(self):
        md = export_markdown(_sample_report())
        assert "What drove the revenue increase?" in md

    def test_issue_heading_with_severity(self):
        md = export_markdown(_sample_report())
        assert "HIGH" in md or "[HIGH]" in md

    def test_advisor_notes_in_code_block(self):
        md = export_markdown(_sample_report())
        assert "```" in md
        assert "ENGAGEMENT NOTES" in md

    def test_materiality_in_header(self):
        md = export_markdown(_sample_report())
        assert "$1,000" in md

    def test_no_issues_state(self):
        report = _sample_report()
        report["triggered_issues"] = []
        md = export_markdown(report)
        assert "No issues detected" in md


# ---------------------------------------------------------------------------
# export_excel
# ---------------------------------------------------------------------------

class TestExportExcel:
    def test_returns_bytes(self):
        result = export_excel(_sample_report())
        assert isinstance(result, bytes)
        assert len(result) > 1000

    def test_valid_xlsx_with_five_sheets(self):
        from openpyxl import load_workbook
        result = export_excel(_sample_report())
        wb = load_workbook(io.BytesIO(result))
        assert "Summary" in wb.sheetnames
        assert "Financial Changes" in wb.sheetnames
        assert "Issues" in wb.sheetnames
        assert "Procedures & AJEs" in wb.sheetnames
        assert "Mgmt Questions" in wb.sheetnames

    def test_summary_sheet_has_entity_name(self):
        from openpyxl import load_workbook
        result = export_excel(_sample_report())
        wb = load_workbook(io.BytesIO(result))
        ws = wb["Summary"]
        all_values = [ws.cell(r, 1).value for r in range(1, 20) if ws.cell(r, 1).value]
        combined = " ".join(str(v) for v in all_values)
        assert "Test Company LLC" in combined

    def test_financial_changes_sheet_has_data(self):
        from openpyxl import load_workbook
        result = export_excel(_sample_report())
        wb = load_workbook(io.BytesIO(result))
        ws = wb["Financial Changes"]
        assert ws.cell(1, 1).value == "Metric"
        assert ws.cell(2, 1).value == "Revenue"

    def test_issues_sheet_has_issue_row(self):
        from openpyxl import load_workbook
        result = export_excel(_sample_report())
        wb = load_workbook(io.BytesIO(result))
        ws = wb["Issues"]
        row2 = [ws.cell(2, c).value for c in range(1, 5)]
        assert "REV_001" in row2

    def test_mgmt_questions_sheet_has_questions(self):
        from openpyxl import load_workbook
        result = export_excel(_sample_report())
        wb = load_workbook(io.BytesIO(result))
        ws = wb["Mgmt Questions"]
        questions = [ws.cell(r, 2).value for r in range(2, 10) if ws.cell(r, 2).value]
        assert any("revenue" in str(q).lower() for q in questions)

    def test_empty_issues_report(self):
        report = _sample_report()
        report["triggered_issues"] = []
        report["suggested_procedures"] = []
        report["suggested_adjustments"] = []
        result = export_excel(report)
        assert isinstance(result, bytes)
        assert len(result) > 1000
