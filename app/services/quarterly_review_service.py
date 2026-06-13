"""
Quarterly Review Generator — Sprint 3.14

Generates a deterministic, structured 8-section quarterly review report by
combining the detection engine, financial diagnostics, and issue repository
template content. No AI generation — all narrative is template-based.

Sections
--------
1. Executive Summary
2. Key Financial Changes
3. Significant Variances
4. Triggered Accounting Issues
5. Suggested Questions for Management
6. Suggested Procedures
7. Suggested Adjustments
8. Advisor Notes
"""
from __future__ import annotations

import datetime
import io
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

from sqlalchemy.orm import Session

from app.models.accounting_period import AccountingPeriod
from app.models.entity import Entity
from app.services import accounting_intelligence_service as svc
from app.services import issue_template_service as tmpl_svc


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

_SEVERITY_ORDER = {"critical": 4, "high": 3, "moderate": 2, "low": 1, "informational": 0}

_METRIC_LABELS: dict[str, str] = {
    "revenue": "Revenue",
    "cogs": "Cost of Goods Sold",
    "gross_profit": "Gross Profit",
    "gross_margin_pct": "Gross Margin %",
    "net_income": "Net Income",
    "total_expenses": "Total Operating Expenses",
    "cash": "Cash & Equivalents",
    "accounts_receivable": "Accounts Receivable",
    "inventory": "Inventory",
    "total_current_assets": "Total Current Assets",
    "total_current_liabilities": "Total Current Liabilities",
    "total_assets": "Total Assets",
    "total_liabilities": "Total Liabilities",
    "total_equity": "Total Equity",
    "total_debt": "Total Debt",
    "working_capital": "Working Capital",
    "current_ratio": "Current Ratio",
    "debt_to_equity": "Debt-to-Equity",
    "roa": "Return on Assets %",
}

_DISPLAY_ORDER = [
    "revenue", "cogs", "gross_profit", "gross_margin_pct",
    "total_expenses", "net_income",
    "cash", "accounts_receivable", "inventory",
    "total_current_assets", "total_current_liabilities",
    "working_capital", "current_ratio",
    "total_assets", "total_liabilities", "total_equity",
    "total_debt", "debt_to_equity", "roa",
]


def _pct_change(current: str | None, prior: str | None) -> Decimal | None:
    try:
        c, p = Decimal(current or "0"), Decimal(prior or "0")
        if p == 0:
            return None
        return ((c - p) / abs(p) * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    except Exception:
        return None


def _fmt_amount(val: str | None) -> str:
    if not val:
        return "—"
    try:
        n = float(val)
    except (ValueError, TypeError):
        return val
    if abs(n) >= 1_000_000:
        return f"${n / 1_000_000:,.2f}M"
    if abs(n) >= 1_000:
        return f"${n / 1_000:,.1f}K"
    return f"${n:,.0f}"


def _fmt_pct(val: Decimal | None) -> str:
    if val is None:
        return "—"
    return f"{val:+.1f}%"


# ---------------------------------------------------------------------------
# Financial change row builder
# ---------------------------------------------------------------------------

def _build_financial_changes(
    cur: dict,
    pri: dict,
    materiality: Decimal,
) -> list[dict]:
    rows = []
    for key in _DISPLAY_ORDER:
        cur_val = cur.get(key)
        pri_val = pri.get(key)
        if cur_val is None and pri_val is None:
            continue

        try:
            c_num = Decimal(cur_val or "0")
            p_num = Decimal(pri_val or "0")
        except Exception:
            continue

        change_amt = c_num - p_num
        change_pct = _pct_change(cur_val, pri_val)

        is_pct_metric = key in ("gross_margin_pct", "roa")
        is_ratio_metric = key in ("current_ratio", "debt_to_equity")

        if is_pct_metric:
            label_current = f"{c_num:.1f}%"
            label_prior = f"{p_num:.1f}%"
            label_change = f"{change_amt:+.1f}pp"
        elif is_ratio_metric:
            label_current = f"{c_num:.2f}x"
            label_prior = f"{p_num:.2f}x"
            label_change = f"{change_amt:+.2f}x"
        else:
            label_current = _fmt_amount(cur_val)
            label_prior = _fmt_amount(pri_val)
            label_change = _fmt_amount(str(change_amt)) if change_amt != 0 else "—"

        direction = (
            "increase" if change_amt > 0
            else "decrease" if change_amt < 0
            else "unchanged"
        )

        # Materiality: amount metrics use threshold; ratio/pct use absolute delta
        if is_pct_metric or is_ratio_metric:
            significant = abs(change_amt) >= Decimal("1")
        else:
            significant = abs(change_amt) >= materiality

        rows.append({
            "metric": key,
            "label": _METRIC_LABELS.get(key, key),
            "current_value": label_current,
            "prior_value": label_prior,
            "change_amount": label_change,
            "change_pct": _fmt_pct(change_pct) if not is_pct_metric and not is_ratio_metric else None,
            "direction": direction,
            "significant": significant,
        })
    return rows


# ---------------------------------------------------------------------------
# Template enrichment
# ---------------------------------------------------------------------------

def _enrich_issues_from_templates(
    db: Session,
    issues: list[dict],
) -> list[dict]:
    """
    For each detected issue, look up repository templates by category to pull
    management_questions, suggested_procedures, and suggested_ajes.

    The detection engine uses codes like AR_GROWTH_EXCEEDS_REVENUE; the
    repository uses AR_003 etc. We match by category.
    """
    tmpl_svc.seed_issue_templates(db)

    # Cache templates by category (top-3 per category, highest risk_level first)
    risk_rank = {"critical": 4, "high": 3, "moderate": 2, "low": 1}
    category_templates: dict[str, list[dict]] = {}

    for issue in issues:
        cat = issue.get("category", "")
        if cat in category_templates:
            continue
        templates = tmpl_svc.list_templates(db, category=cat)
        templates.sort(key=lambda t: risk_rank.get(t.get("risk_level", "low"), 0), reverse=True)
        category_templates[cat] = templates[:3]

    enriched = []
    for issue in issues:
        cat = issue.get("category", "")
        templates = category_templates.get(cat, [])

        mgmt_qs: list[str] = []
        procedures: list[str] = []
        ajes: list[str] = []

        for tmpl in templates:
            mgmt_qs.extend(tmpl.get("management_questions", []))
            procedures.extend(tmpl.get("suggested_procedures", []))
            ajes.extend(tmpl.get("suggested_ajes", []))

        # Deduplicate while preserving order
        seen: set[str] = set()
        unique_qs: list[str] = []
        for q in mgmt_qs:
            if q not in seen:
                seen.add(q)
                unique_qs.append(q)

        # Merge with inline procedures/AJEs from detection result
        inline_procedures = [
            p.strip()
            for p in (issue.get("suggested_procedures") or "").split("\n")
            if p.strip() and not p.strip().startswith("#")
        ]
        inline_ajes = [
            a.strip()
            for a in (issue.get("suggested_ajes") or "").split("\n")
            if a.strip() and not a.strip().startswith("#")
        ]

        enriched.append({
            **issue,
            "management_questions": unique_qs[:5],
            "procedure_items": inline_procedures or [p for p in procedures[:5]],
            "aje_items": inline_ajes or [a for a in ajes[:3]],
        })

    return enriched


# ---------------------------------------------------------------------------
# Deterministic narrative builders
# ---------------------------------------------------------------------------

def _build_key_findings(
    financial_changes: list[dict],
    issues: list[dict],
    cur_diag: dict,
    pri_diag: dict,
) -> list[str]:
    findings = []

    # Revenue trend
    rev_row = next((r for r in financial_changes if r["metric"] == "revenue"), None)
    if rev_row and rev_row["direction"] != "unchanged" and rev_row["change_pct"]:
        direction_word = "increased" if rev_row["direction"] == "increase" else "declined"
        findings.append(
            f"Revenue {direction_word} {rev_row['change_pct']} period-over-period "
            f"({rev_row['prior_value']} → {rev_row['current_value']})."
        )

    # Gross margin
    gm_row = next((r for r in financial_changes if r["metric"] == "gross_margin_pct"), None)
    if gm_row and gm_row["significant"]:
        direction_word = "improved" if gm_row["direction"] == "increase" else "compressed"
        findings.append(
            f"Gross margin {direction_word} from {gm_row['prior_value']} to {gm_row['current_value']}."
        )

    # Net income
    ni_row = next((r for r in financial_changes if r["metric"] == "net_income"), None)
    if ni_row and ni_row["significant"]:
        direction_word = "increased" if ni_row["direction"] == "increase" else "decreased"
        findings.append(f"Net income {direction_word} to {ni_row['current_value']}.")

    # Issue summary
    critical = [i for i in issues if i.get("severity") == "critical"]
    high = [i for i in issues if i.get("severity") == "high"]
    if critical:
        findings.append(
            f"{len(critical)} critical-severity issue{'s' if len(critical) > 1 else ''} detected "
            f"requiring immediate review: {', '.join(i['title'] for i in critical[:2])}."
        )
    elif high:
        findings.append(
            f"{len(high)} high-severity issue{'s' if len(high) > 1 else ''} identified: "
            f"{', '.join(i['title'] for i in high[:2])}."
        )
    elif issues:
        findings.append(f"{len(issues)} accounting issue{'s' if len(issues) > 1 else ''} detected at moderate or lower severity.")
    else:
        findings.append("No accounting issues detected above the materiality threshold.")

    return findings[:5]


def _build_executive_narrative(
    entity_name: str,
    current_period_name: str,
    comparison_period_name: str,
    issues: list[dict],
    financial_changes: list[dict],
) -> str:
    issue_count = len(issues)
    sev_counts = {s: sum(1 for i in issues if i.get("severity") == s)
                  for s in ("critical", "high", "moderate", "low")}

    risk = (
        "elevated" if sev_counts["critical"] > 0 or sev_counts["high"] > 1
        else "moderate" if sev_counts["high"] > 0 or sev_counts["moderate"] > 2
        else "low"
    )

    rev = next((r for r in financial_changes if r["metric"] == "revenue"), None)
    rev_clause = (
        f" Revenue {('increased' if rev['direction'] == 'increase' else 'declined')} "
        f"{rev['change_pct']} compared to {comparison_period_name}."
        if rev and rev.get("change_pct") else ""
    )

    issue_clause = (
        f" {issue_count} accounting {'issue was' if issue_count == 1 else 'issues were'} "
        f"identified by the detection engine"
        + (
            f", including {sev_counts['critical']} critical and {sev_counts['high']} high-severity items"
            if sev_counts["critical"] or sev_counts["high"] else
            f" at moderate or lower severity"
        )
        + "."
    ) if issue_count > 0 else " No accounting issues were detected above the materiality threshold."

    return (
        f"This automated review covers {entity_name}'s financial position for {current_period_name}, "
        f"compared to {comparison_period_name}.{rev_clause}{issue_clause} "
        f"Overall risk assessment: {risk.upper()}. "
        f"All findings are deterministic rule-based detections and should be evaluated "
        f"in the context of the full engagement scope."
    )


def _build_advisor_notes(
    entity_name: str,
    current_period_name: str,
    comparison_period_name: str,
    issues: list[dict],
    financial_changes: list[dict],
    materiality: Decimal,
) -> str:
    open_issues = [i for i in issues if i.get("status") == "open"]
    sev_counts = {s: sum(1 for i in open_issues if i.get("severity") == s)
                  for s in ("critical", "high", "moderate", "low")}

    high_priority = [i for i in open_issues if i.get("severity") in ("critical", "high")]
    hp_titles = "; ".join(i["title"] for i in high_priority[:3])

    ni_row = next((r for r in financial_changes if r["metric"] == "net_income"), None)
    cr_row = next((r for r in financial_changes if r["metric"] == "current_ratio"), None)

    sections = [
        f"ENGAGEMENT NOTES — {entity_name.upper()} — {current_period_name.upper()}",
        "",
        f"Period under review:     {current_period_name}",
        f"Comparison period:       {comparison_period_name}",
        f"Materiality threshold:   ${materiality:,.0f}",
        f"Open issues:             {len(open_issues)} "
        f"(critical: {sev_counts['critical']}, high: {sev_counts['high']}, "
        f"moderate: {sev_counts['moderate']}, low: {sev_counts['low']})",
        "",
    ]

    if high_priority:
        sections += [
            "PRIORITY ITEMS REQUIRING ATTENTION:",
            f"  {hp_titles}",
            "",
            "Recommend discussion with management before issuing any report or representation.",
            "",
        ]

    if ni_row and ni_row["significant"]:
        direction = "improvement" if ni_row["direction"] == "increase" else "decline"
        sections.append(
            f"Net income showed a {direction} to {ni_row['current_value']}. "
            f"Corroborate with supporting schedules."
        )

    if cr_row and cr_row["significant"]:
        direction = "improved" if cr_row["direction"] == "increase" else "declined"
        sections.append(
            f"Liquidity (current ratio) {direction} from {cr_row['prior_value']} to {cr_row['current_value']}."
        )

    sections += [
        "",
        "STANDARD NOTES:",
        "  • All findings are machine-generated and require professional review.",
        "  • Suggested procedures are guidelines; scope them per engagement risk.",
        "  • Suggested AJEs require supporting documentation before posting.",
        "  • Management questions should be documented in workpapers.",
    ]

    return "\n".join(sections)


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def generate_review(
    db: Session,
    entity_id: int,
    current_period_id: int,
    comparison_period_id: int,
    scenario_id: int | None = None,
    materiality_threshold: Decimal = Decimal("1000"),
) -> dict:
    """
    Generate a complete quarterly review report.

    Runs detection, computes diagnostics for both periods, enriches with
    repository template content, and assembles 8 report sections.
    """
    entity = db.get(Entity, entity_id)
    if entity is None:
        raise ValueError(f"Entity {entity_id} not found")

    cur_period = db.get(AccountingPeriod, current_period_id)
    cmp_period = db.get(AccountingPeriod, comparison_period_id)
    if cur_period is None:
        raise ValueError(f"Period {current_period_id} not found")
    if cmp_period is None:
        raise ValueError(f"Period {comparison_period_id} not found")

    # 1. Run detection and collect issues
    issues_raw = svc.run_detection(
        db,
        entity_id=entity_id,
        current_period_id=current_period_id,
        comparison_period_id=comparison_period_id,
        scenario_id=scenario_id,
        materiality_threshold=materiality_threshold,
        persist=True,
    )

    # 2. Compute diagnostics for both periods
    cur_diag = svc.compute_diagnostics(db, entity_id, current_period_id, scenario_id)
    cmp_diag = svc.compute_diagnostics(db, entity_id, comparison_period_id, scenario_id)

    # 3. Build financial change rows
    financial_changes = _build_financial_changes(cur_diag, cmp_diag, materiality_threshold)
    significant_variances = [r for r in financial_changes if r["significant"]]

    # 4. Enrich issues with template content
    issues = _enrich_issues_from_templates(db, issues_raw)

    # Sort by severity descending
    issues.sort(key=lambda i: _SEVERITY_ORDER.get(i.get("severity", "low"), 0), reverse=True)

    # 5. Build aggregated sections
    severity_counts = {
        s: sum(1 for i in issues if i.get("severity") == s)
        for s in ("critical", "high", "moderate", "low", "informational")
    }
    overall_risk = (
        "elevated" if severity_counts["critical"] > 0 or severity_counts["high"] > 1
        else "moderate" if severity_counts["high"] > 0 or severity_counts["moderate"] > 2
        else "low"
    )
    key_findings = _build_key_findings(financial_changes, issues, cur_diag, cmp_diag)

    # Aggregate all management questions (unique, capped at 20)
    seen: set[str] = set()
    all_mgmt_qs: list[str] = []
    for issue in issues:
        for q in issue.get("management_questions", []):
            if q not in seen:
                seen.add(q)
                all_mgmt_qs.append(q)
    all_mgmt_qs = all_mgmt_qs[:20]

    suggested_procedures = [
        {
            "issue_code": i["issue_code"],
            "title": i["title"],
            "severity": i.get("severity", "moderate"),
            "items": i.get("procedure_items", []),
        }
        for i in issues
        if i.get("procedure_items")
    ]

    suggested_adjustments = [
        {
            "issue_code": i["issue_code"],
            "title": i["title"],
            "severity": i.get("severity", "moderate"),
            "items": i.get("aje_items", []),
        }
        for i in issues
        if i.get("aje_items")
    ]

    exec_narrative = _build_executive_narrative(
        entity.name,
        cur_period.period_name,
        cmp_period.period_name,
        issues,
        financial_changes,
    )

    advisor_notes = _build_advisor_notes(
        entity.name,
        cur_period.period_name,
        cmp_period.period_name,
        issues,
        financial_changes,
        materiality_threshold,
    )

    # 6. Assemble report
    report: dict[str, Any] = {
        "metadata": {
            "entity_id": entity_id,
            "entity_name": entity.name,
            "current_period": {
                "id": cur_period.id,
                "name": cur_period.period_name,
                "start": str(cur_period.start_date),
                "end": str(cur_period.end_date),
            },
            "comparison_period": {
                "id": cmp_period.id,
                "name": cmp_period.period_name,
                "start": str(cmp_period.start_date),
                "end": str(cmp_period.end_date),
            },
            "generated_at": datetime.datetime.utcnow().isoformat() + "Z",
            "materiality_threshold": str(materiality_threshold),
        },
        "executive_summary": {
            "overall_risk": overall_risk,
            "total_issues": len(issues),
            "issues_by_severity": severity_counts,
            "key_findings": key_findings,
            "narrative": exec_narrative,
        },
        "key_financial_changes": financial_changes,
        "significant_variances": significant_variances,
        "triggered_issues": [
            {
                "issue_code": i["issue_code"],
                "category": i["category"],
                "severity": i["severity"],
                "title": i["title"],
                "description": i["description"],
                "detection_trigger": i.get("detection_trigger"),
                "supporting_metrics": i.get("supporting_metrics", {}),
                "management_questions": i.get("management_questions", []),
                "procedure_items": i.get("procedure_items", []),
                "aje_items": i.get("aje_items", []),
                "status": i.get("status", "open"),
            }
            for i in issues
        ],
        "management_questions": all_mgmt_qs,
        "suggested_procedures": suggested_procedures,
        "suggested_adjustments": suggested_adjustments,
        "advisor_notes": advisor_notes,
    }
    return report


# ---------------------------------------------------------------------------
# Markdown export
# ---------------------------------------------------------------------------

def export_markdown(report: dict) -> str:
    md: list[str] = []
    meta = report["metadata"]
    summ = report["executive_summary"]

    md.append(f"# Quarterly Review — {meta['entity_name']}")
    md.append(f"**Period:** {meta['current_period']['name']}  ")
    md.append(f"**Compared to:** {meta['comparison_period']['name']}  ")
    md.append(f"**Generated:** {meta['generated_at'][:10]}  ")
    md.append(f"**Materiality:** ${float(meta['materiality_threshold']):,.0f}")
    md.append("")

    # Section 1
    md.append("---")
    md.append(f"## 1. Executive Summary")
    md.append(f"**Overall Risk:** {summ['overall_risk'].upper()}  ")
    md.append(f"**Issues Detected:** {summ['total_issues']}  ")
    sevs = summ["issues_by_severity"]
    md.append(
        f"*(Critical: {sevs['critical']} · High: {sevs['high']} · "
        f"Moderate: {sevs['moderate']} · Low: {sevs['low']})*"
    )
    md.append("")
    md.append(summ["narrative"])
    md.append("")
    if summ["key_findings"]:
        md.append("### Key Findings")
        for f in summ["key_findings"]:
            md.append(f"- {f}")
        md.append("")

    # Section 2
    md.append("---")
    md.append("## 2. Key Financial Changes")
    md.append("| Metric | Prior | Current | Change |")
    md.append("|--------|-------|---------|--------|")
    for row in report["key_financial_changes"]:
        chg = row["change_amount"]
        if row.get("change_pct"):
            chg += f" ({row['change_pct']})"
        flag = " ⚠" if row.get("significant") else ""
        md.append(f"| {row['label']}{flag} | {row['prior_value']} | {row['current_value']} | {chg} |")
    md.append("")

    # Section 3
    md.append("---")
    md.append("## 3. Significant Variances")
    if report["significant_variances"]:
        for row in report["significant_variances"]:
            md.append(f"- **{row['label']}**: {row['prior_value']} → {row['current_value']} ({row['change_amount']})")
    else:
        md.append("*No variances exceeded the materiality threshold.*")
    md.append("")

    # Section 4
    md.append("---")
    md.append("## 4. Triggered Accounting Issues")
    for issue in report["triggered_issues"]:
        sev_badge = issue["severity"].upper()
        md.append(f"### [{sev_badge}] {issue['title']}")
        md.append(f"*Code: {issue['issue_code']} · Category: {issue['category']}*")
        md.append("")
        md.append(issue["description"])
        if issue.get("detection_trigger"):
            md.append(f"**Trigger:** {issue['detection_trigger']}")
        md.append("")
    if not report["triggered_issues"]:
        md.append("*No issues detected.*")
        md.append("")

    # Section 5
    md.append("---")
    md.append("## 5. Suggested Questions for Management")
    for i, q in enumerate(report["management_questions"], 1):
        md.append(f"{i}. {q}")
    if not report["management_questions"]:
        md.append("*No management questions generated.*")
    md.append("")

    # Section 6
    md.append("---")
    md.append("## 6. Suggested Procedures")
    for proc in report["suggested_procedures"]:
        md.append(f"### {proc['title']}")
        for item in proc["items"]:
            md.append(f"- {item}")
        md.append("")
    if not report["suggested_procedures"]:
        md.append("*No suggested procedures.*")
        md.append("")

    # Section 7
    md.append("---")
    md.append("## 7. Suggested Adjusting Journal Entries")
    for adj in report["suggested_adjustments"]:
        md.append(f"### {adj['title']}")
        for item in adj["items"]:
            md.append(f"- {item}")
        md.append("")
    if not report["suggested_adjustments"]:
        md.append("*No suggested adjustments.*")
        md.append("")

    # Section 8
    md.append("---")
    md.append("## 8. Advisor Notes")
    md.append("```")
    md.append(report["advisor_notes"])
    md.append("```")

    return "\n".join(md)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_excel(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    meta = report["metadata"]
    summ = report["executive_summary"]

    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
    ALT_FILL = PatternFill("solid", fgColor="DCE6F1")
    WARN_FILL = PatternFill("solid", fgColor="FFE699")
    WHITE_FONT = Font(bold=True, color="FFFFFF")
    BOLD = Font(bold=True)

    def _header(ws, row: int, col: int, text: str, fill=HEADER_FILL):
        c = ws.cell(row=row, column=col, value=text)
        c.font = WHITE_FONT
        c.fill = fill
        c.alignment = Alignment(wrap_text=True)
        return c

    wb = Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"

    ws.cell(1, 1, f"Quarterly Review — {meta['entity_name']}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Period: {meta['current_period']['name']}")
    ws.cell(3, 1, f"Compared to: {meta['comparison_period']['name']}")
    ws.cell(4, 1, f"Generated: {meta['generated_at'][:10]}")
    ws.cell(5, 1, f"Materiality: ${float(meta['materiality_threshold']):,.0f}")
    ws.cell(6, 1, f"Overall Risk: {summ['overall_risk'].upper()}").font = Font(bold=True)
    ws.cell(7, 1, f"Total Issues: {summ['total_issues']}")

    r = 9
    ws.cell(r, 1, "Key Findings").font = BOLD
    for finding in summ["key_findings"]:
        r += 1
        ws.cell(r, 1, f"• {finding}")

    r += 2
    ws.cell(r, 1, "Executive Summary").font = BOLD
    r += 1
    ws.cell(r, 1, summ["narrative"]).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 80
    ws.column_dimensions["A"].width = 100

    # ── Sheet 2: Financial Changes ────────────────────────────────────────────
    ws2 = wb.create_sheet("Financial Changes")
    headers = ["Metric", "Prior", "Current", "Change", "Change %", "Material?"]
    for ci, h in enumerate(headers, 1):
        _header(ws2, 1, ci, h)
    ws2.column_dimensions["A"].width = 30
    for ci in range(2, 7):
        ws2.column_dimensions[get_column_letter(ci)].width = 16

    for ri, row in enumerate(report["key_financial_changes"], 2):
        fill = WARN_FILL if row.get("significant") else (ALT_FILL if ri % 2 == 0 else None)
        cells = [
            row["label"],
            row["prior_value"],
            row["current_value"],
            row["change_amount"],
            row.get("change_pct") or "—",
            "Yes" if row.get("significant") else "",
        ]
        for ci, val in enumerate(cells, 1):
            c = ws2.cell(ri, ci, val)
            if fill:
                c.fill = fill

    # ── Sheet 3: Triggered Issues ─────────────────────────────────────────────
    ws3 = wb.create_sheet("Issues")
    issue_headers = ["Code", "Severity", "Category", "Title", "Description", "Trigger"]
    for ci, h in enumerate(issue_headers, 1):
        _header(ws3, 1, ci, h)
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 20
    ws3.column_dimensions["D"].width = 40
    ws3.column_dimensions["E"].width = 60
    ws3.column_dimensions["F"].width = 50

    SEV_FILLS = {
        "critical": PatternFill("solid", fgColor="FF0000"),
        "high": PatternFill("solid", fgColor="FF6600"),
        "moderate": PatternFill("solid", fgColor="FFC000"),
        "low": PatternFill("solid", fgColor="92D050"),
    }
    for ri, issue in enumerate(report["triggered_issues"], 2):
        row_data = [
            issue["issue_code"], issue["severity"], issue["category"],
            issue["title"], issue["description"],
            issue.get("detection_trigger") or "",
        ]
        for ci, val in enumerate(row_data, 1):
            c = ws3.cell(ri, ci, val)
            c.alignment = Alignment(wrap_text=True)
            if ci == 2:
                c.fill = SEV_FILLS.get(issue["severity"], PatternFill())

    # ── Sheet 4: Procedures & AJEs ────────────────────────────────────────────
    ws4 = wb.create_sheet("Procedures & AJEs")
    _header(ws4, 1, 1, "Issue")
    _header(ws4, 1, 2, "Type")
    _header(ws4, 1, 3, "Item")
    ws4.column_dimensions["A"].width = 40
    ws4.column_dimensions["B"].width = 14
    ws4.column_dimensions["C"].width = 80

    r4 = 2
    for proc in report["suggested_procedures"]:
        for item in proc["items"]:
            ws4.cell(r4, 1, proc["title"])
            ws4.cell(r4, 2, "Procedure")
            ws4.cell(r4, 3, item).alignment = Alignment(wrap_text=True)
            r4 += 1
    for adj in report["suggested_adjustments"]:
        for item in adj["items"]:
            ws4.cell(r4, 1, adj["title"])
            ws4.cell(r4, 2, "AJE")
            ws4.cell(r4, 3, item).alignment = Alignment(wrap_text=True)
            r4 += 1

    # ── Sheet 5: Management Questions ─────────────────────────────────────────
    ws5 = wb.create_sheet("Mgmt Questions")
    _header(ws5, 1, 1, "#")
    _header(ws5, 1, 2, "Question")
    ws5.column_dimensions["A"].width = 6
    ws5.column_dimensions["B"].width = 100
    for ri, q in enumerate(report["management_questions"], 2):
        ws5.cell(ri, 1, ri - 1)
        ws5.cell(ri, 2, q).alignment = Alignment(wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
