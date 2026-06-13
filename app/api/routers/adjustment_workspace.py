"""
Adjustment Workspace API — unified adjustment management for advisors.

Provides: adjustment listing with impact, materiality, packages, advisor notes,
multi-select impact preview, account-level rollforward, and advisor scenarios.
"""
from __future__ import annotations

import datetime
from decimal import Decimal
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.api.deps import get_db, get_current_user
from app.api.schemas import (
    AdvisorNoteOut,
    AdvisorNoteUpdate,
    AdvisorScenarioCreate,
    AdvisorScenarioOut,
    AdvisorScenarioPackageIn,
    AdvisorScenarioPackageItem,
    AdvisorScenarioUpdate,
    ADVISOR_SCENARIO_TYPES,
    AdjustmentImpact,
    AdjustmentListItem,
    AdjustmentPackageCreate,
    AdjustmentPackageOut,
    AdjustmentPackageUpdate,
    ImpactPreviewRequest,
    MaterialityUpdate,
    PackageToggleIn,
    PACKAGE_TYPES_EXTENDED,
    RollforwardRow,
    ScenarioComparisonResult,
    ScenarioImpactResult,
)
from app.models.account import Account
from app.models.adjustment_workspace import (
    AdjustmentAdvisorNote,
    AdjustmentPackage,
    AdjustmentPackageMembership,
)
from app.models.advisor_scenario import AdvisorScenario, AdvisorScenarioPackage
from app.models.journal_entry import JournalEntry
from app.models.journal_entry_line import JournalEntryLine

router = APIRouter(prefix="/adjustment-workspace", tags=["adjustment-workspace"])

_INCOME_TYPES = {"revenue", "other_income"}
_EXPENSE_TYPES = {"cogs", "expense", "other_expense", "tax"}
_ASSET_TYPES = {"asset"}
_LIABILITY_TYPES = {"liability", "intercompany"}
_EQUITY_TYPES = {"equity"}


def _compute_impact(db: Session, je_id: int) -> AdjustmentImpact:
    lines = (
        db.query(JournalEntryLine, Account)
        .join(Account, JournalEntryLine.account_id == Account.id)
        .filter(JournalEntryLine.journal_entry_id == je_id)
        .all()
    )
    ni = asset = liability = equity = Decimal("0")
    for line, acct in lines:
        net = Decimal(str(line.debit)) - Decimal(str(line.credit))
        t = acct.account_type
        if t in _INCOME_TYPES:
            ni -= net
        elif t in _EXPENSE_TYPES:
            ni += net
        if t in _ASSET_TYPES:
            asset += net
        elif t in _LIABILITY_TYPES:
            liability += net
        elif t in _EQUITY_TYPES:
            equity += net
    return AdjustmentImpact(
        ni_impact=ni,
        ebitda_impact=ni,
        asset_impact=asset,
        liability_impact=liability,
        equity_impact=equity,
    )


def _total_debit(db: Session, je_id: int) -> Decimal:
    result = (
        db.query(func.sum(JournalEntryLine.debit))
        .filter(JournalEntryLine.journal_entry_id == je_id)
        .scalar()
    )
    return Decimal(str(result or 0))


def _package_ids_for_je(db: Session, je_id: int) -> list[int]:
    rows = (
        db.query(AdjustmentPackageMembership.package_id)
        .filter(AdjustmentPackageMembership.journal_entry_id == je_id)
        .all()
    )
    return [r[0] for r in rows]


def _note_for_je(db: Session, je_id: int) -> AdjustmentAdvisorNote | None:
    return (
        db.query(AdjustmentAdvisorNote)
        .filter(AdjustmentAdvisorNote.journal_entry_id == je_id)
        .first()
    )


def _org_from_user(user) -> str:
    if user is None:
        return "default-org"
    org = getattr(user, "organization_id", None)
    return str(org) if org is not None else "default-org"


# ---------------------------------------------------------------------------
# Adjustment listing
# ---------------------------------------------------------------------------

@router.get("/adjustments", response_model=list[AdjustmentListItem])
def list_adjustments(
    entity_id: Optional[int] = Query(None),
    scenario_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    overlay_group: Optional[str] = Query(None),
    materiality: Optional[str] = Query(None),
    package_id: Optional[int] = Query(None),
    search: Optional[str] = Query(None),
    date_from: Optional[datetime.date] = Query(None),
    date_to: Optional[datetime.date] = Query(None),
    limit: int = Query(200, le=1000),
    offset: int = Query(0),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    q = db.query(JournalEntry)

    if entity_id is not None:
        q = q.filter(JournalEntry.entity_id == entity_id)
    if scenario_id is not None:
        q = q.filter(JournalEntry.scenario_id == scenario_id)
    if status is not None:
        q = q.filter(JournalEntry.status == status)
    if overlay_group is not None:
        q = q.filter(JournalEntry.overlay_group == overlay_group)
    if materiality is not None:
        q = q.filter(JournalEntry.materiality == materiality)
    if search:
        like = f"%{search}%"
        q = q.filter(
            JournalEntry.description.ilike(like) | JournalEntry.je_number.ilike(like)
        )
    if date_from is not None:
        q = q.filter(JournalEntry.entry_date >= date_from)
    if date_to is not None:
        q = q.filter(JournalEntry.entry_date <= date_to)
    if package_id is not None:
        member_ids = (
            db.query(AdjustmentPackageMembership.journal_entry_id)
            .filter(AdjustmentPackageMembership.package_id == package_id)
            .subquery()
        )
        q = q.filter(JournalEntry.id.in_(member_ids))

    jes = q.order_by(JournalEntry.entry_date.desc(), JournalEntry.id.desc()).offset(offset).limit(limit).all()

    items = []
    for je in jes:
        note = _note_for_je(db, je.id)
        items.append(
            AdjustmentListItem(
                id=je.id,
                je_number=je.je_number,
                entry_date=je.entry_date,
                entity_id=je.entity_id,
                scenario_id=je.scenario_id,
                description=je.description,
                source=je.source,
                status=je.status,
                overlay_group=je.overlay_group,
                materiality=je.materiality,
                total_debit=_total_debit(db, je.id),
                impact=_compute_impact(db, je.id),
                package_ids=_package_ids_for_je(db, je.id),
                has_advisor_note=note is not None,
                advisor_resolution_status=note.resolution_status if note else None,
            )
        )
    return items


@router.patch("/adjustments/{je_id}/materiality", response_model=dict)
def set_materiality(
    je_id: int,
    body: MaterialityUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    je = db.get(JournalEntry, je_id)
    if je is None:
        raise HTTPException(status_code=404, detail="Journal entry not found")
    valid = {None, "clearly_trivial", "immaterial", "material", "critical"}
    if body.materiality not in valid:
        raise HTTPException(status_code=422, detail=f"Invalid materiality value")
    je.materiality = body.materiality
    je.updated_at = datetime.datetime.utcnow()
    db.flush()
    return {"id": je_id, "materiality": je.materiality}


# ---------------------------------------------------------------------------
# Packages
# ---------------------------------------------------------------------------

@router.get("/packages", response_model=list[AdjustmentPackageOut])
def list_packages(
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org = _org_from_user(user)
    pkgs = (
        db.query(AdjustmentPackage)
        .filter(AdjustmentPackage.organization_id == org)
        .order_by(AdjustmentPackage.created_at.desc())
        .all()
    )
    result = []
    for pkg in pkgs:
        count = (
            db.query(func.count(AdjustmentPackageMembership.id))
            .filter(AdjustmentPackageMembership.package_id == pkg.id)
            .scalar()
        ) or 0
        result.append(
            AdjustmentPackageOut(
                id=pkg.id,
                organization_id=pkg.organization_id,
                name=pkg.name,
                package_type=pkg.package_type,
                status=pkg.status,
                description=pkg.description,
                created_at=pkg.created_at,
                updated_at=pkg.updated_at,
                member_count=count,
            )
        )
    return result


@router.post("/packages", response_model=AdjustmentPackageOut, status_code=201)
def create_package(
    body: AdjustmentPackageCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org = _org_from_user(user)
    if body.package_type not in PACKAGE_TYPES_EXTENDED:
        raise HTTPException(status_code=422, detail="Invalid package_type")
    pkg = AdjustmentPackage(
        organization_id=org,
        name=body.name,
        package_type=body.package_type,
        description=body.description,
        created_by_user_id=getattr(user, "id", None),
    )
    db.add(pkg)
    db.flush()
    return AdjustmentPackageOut(
        id=pkg.id,
        organization_id=pkg.organization_id,
        name=pkg.name,
        package_type=pkg.package_type,
        status=pkg.status,
        description=pkg.description,
        created_at=pkg.created_at,
        updated_at=pkg.updated_at,
        member_count=0,
    )


@router.put("/packages/{pkg_id}", response_model=AdjustmentPackageOut)
def update_package(
    pkg_id: int,
    body: AdjustmentPackageUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    pkg = db.get(AdjustmentPackage, pkg_id)
    if pkg is None:
        raise HTTPException(status_code=404, detail="Package not found")
    if body.name is not None:
        pkg.name = body.name
    if body.package_type is not None:
        pkg.package_type = body.package_type
    if body.status is not None:
        pkg.status = body.status
    if body.description is not None:
        pkg.description = body.description
    pkg.updated_at = datetime.datetime.utcnow()
    db.flush()
    count = (
        db.query(func.count(AdjustmentPackageMembership.id))
        .filter(AdjustmentPackageMembership.package_id == pkg.id)
        .scalar()
    ) or 0
    return AdjustmentPackageOut(
        id=pkg.id,
        organization_id=pkg.organization_id,
        name=pkg.name,
        package_type=pkg.package_type,
        status=pkg.status,
        description=pkg.description,
        created_at=pkg.created_at,
        updated_at=pkg.updated_at,
        member_count=count,
    )


@router.delete("/packages/{pkg_id}", status_code=204)
def delete_package(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    pkg = db.get(AdjustmentPackage, pkg_id)
    if pkg is None:
        raise HTTPException(status_code=404, detail="Package not found")
    db.delete(pkg)


@router.post("/packages/{pkg_id}/members", response_model=dict, status_code=201)
def add_member(
    pkg_id: int,
    body: dict,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    pkg = db.get(AdjustmentPackage, pkg_id)
    if pkg is None:
        raise HTTPException(status_code=404, detail="Package not found")
    je_id = body.get("journal_entry_id")
    if je_id is None:
        raise HTTPException(status_code=422, detail="journal_entry_id required")
    if db.get(JournalEntry, je_id) is None:
        raise HTTPException(status_code=404, detail="Journal entry not found")
    existing = (
        db.query(AdjustmentPackageMembership)
        .filter_by(package_id=pkg_id, journal_entry_id=je_id)
        .first()
    )
    if existing:
        return {"package_id": pkg_id, "journal_entry_id": je_id, "already_member": True}
    member = AdjustmentPackageMembership(
        package_id=pkg_id,
        journal_entry_id=je_id,
        added_by_user_id=getattr(user, "id", None),
    )
    db.add(member)
    db.flush()
    return {"package_id": pkg_id, "journal_entry_id": je_id, "added": True}


@router.delete("/packages/{pkg_id}/members/{je_id}", status_code=204)
def remove_member(
    pkg_id: int,
    je_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    member = (
        db.query(AdjustmentPackageMembership)
        .filter_by(package_id=pkg_id, journal_entry_id=je_id)
        .first()
    )
    if member is None:
        raise HTTPException(status_code=404, detail="Membership not found")
    db.delete(member)


# ---------------------------------------------------------------------------
# Advisor notes
# ---------------------------------------------------------------------------

@router.get("/adjustments/{je_id}/notes", response_model=AdvisorNoteOut | None)
def get_notes(
    je_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    return _note_for_je(db, je_id)


@router.put("/adjustments/{je_id}/notes", response_model=AdvisorNoteOut)
def upsert_notes(
    je_id: int,
    body: AdvisorNoteUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    if db.get(JournalEntry, je_id) is None:
        raise HTTPException(status_code=404, detail="Journal entry not found")
    note = _note_for_je(db, je_id)
    if note is None:
        note = AdjustmentAdvisorNote(journal_entry_id=je_id)
        db.add(note)
    if body.issue is not None:
        note.issue = body.issue
    if body.recommendation is not None:
        note.recommendation = body.recommendation
    if body.client_response is not None:
        note.client_response = body.client_response
    if body.resolution_status is not None:
        note.resolution_status = body.resolution_status
    note.updated_at = datetime.datetime.utcnow()
    note.updated_by_user_id = getattr(user, "id", None)
    db.flush()
    return note


# ---------------------------------------------------------------------------
# Multi-select impact preview
# ---------------------------------------------------------------------------

@router.post("/impact-preview", response_model=AdjustmentImpact)
def impact_preview(
    body: ImpactPreviewRequest,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    ni = asset = liability = equity = Decimal("0")
    if not body.journal_entry_ids:
        return AdjustmentImpact()
    lines = (
        db.query(JournalEntryLine, Account)
        .join(Account, JournalEntryLine.account_id == Account.id)
        .filter(JournalEntryLine.journal_entry_id.in_(body.journal_entry_ids))
        .all()
    )
    for line, acct in lines:
        net = Decimal(str(line.debit)) - Decimal(str(line.credit))
        t = acct.account_type
        if t in _INCOME_TYPES:
            ni -= net
        elif t in _EXPENSE_TYPES:
            ni += net
        if t in _ASSET_TYPES:
            asset += net
        elif t in _LIABILITY_TYPES:
            liability += net
        elif t in _EQUITY_TYPES:
            equity += net
    return AdjustmentImpact(
        ni_impact=ni,
        ebitda_impact=ni,
        asset_impact=asset,
        liability_impact=liability,
        equity_impact=equity,
    )


# ---------------------------------------------------------------------------
# Rollforward
# ---------------------------------------------------------------------------

@router.get("/rollforward", response_model=list[RollforwardRow])
def rollforward(
    entity_id: int = Query(...),
    scenario_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    accts = (
        db.query(Account)
        .filter(Account.entity_id == entity_id, Account.is_postable == True)
        .order_by(Account.account_number)
        .all()
    )

    rows = []
    for acct in accts:
        base_q = (
            db.query(
                func.sum(JournalEntryLine.debit - JournalEntryLine.credit)
            )
            .join(JournalEntry, JournalEntryLine.journal_entry_id == JournalEntry.id)
            .filter(
                JournalEntryLine.account_id == acct.id,
                JournalEntry.status == "posted",
                JournalEntry.scenario_id.in_(
                    db.query(JournalEntry.scenario_id)
                    .filter(JournalEntry.entity_id == entity_id)
                    .filter(JournalEntry.source != "manual")
                    .subquery()
                ) if scenario_id is None else JournalEntry.scenario_id != scenario_id,
            )
        )
        adj_q = (
            db.query(
                func.sum(JournalEntryLine.debit - JournalEntryLine.credit)
            )
            .join(JournalEntry, JournalEntryLine.journal_entry_id == JournalEntry.id)
            .filter(
                JournalEntryLine.account_id == acct.id,
                JournalEntry.entity_id == entity_id,
            )
        )
        if scenario_id is not None:
            adj_q = adj_q.filter(JournalEntry.scenario_id == scenario_id)

        as_reported = Decimal(str(base_q.scalar() or 0))
        adjustments = Decimal(str(adj_q.scalar() or 0))

        if as_reported == 0 and adjustments == 0:
            continue

        rows.append(
            RollforwardRow(
                account_id=acct.id,
                account_number=acct.account_number,
                account_name=acct.name,
                account_type=acct.account_type,
                as_reported=as_reported,
                adjustments=adjustments,
                adjusted=as_reported + adjustments,
            )
        )
    return rows

# ---------------------------------------------------------------------------
# Sprint 3.15 — Advisor Scenarios
# ---------------------------------------------------------------------------

def _scenario_to_out(db: Session, scen: AdvisorScenario) -> AdvisorScenarioOut:
    items: list[AdvisorScenarioPackageItem] = []
    for asp in scen.packages:
        pkg = db.get(AdjustmentPackage, asp.package_id)
        if pkg:
            items.append(AdvisorScenarioPackageItem(
                id=asp.id,
                package_id=asp.package_id,
                package_name=pkg.name,
                package_type=pkg.package_type,
                included=asp.included,
                include_order=asp.include_order,
            ))
    return AdvisorScenarioOut(
        id=scen.id,
        organization_id=scen.organization_id,
        name=scen.name,
        scenario_type=scen.scenario_type,
        description=scen.description,
        created_at=scen.created_at,
        updated_at=scen.updated_at,
        packages=items,
    )


def _compute_impact_for_jes(db: Session, je_ids: list[int]) -> AdjustmentImpact:
    if not je_ids:
        return AdjustmentImpact()
    ni = asset = liability = equity = Decimal("0")
    lines = (
        db.query(JournalEntryLine, Account)
        .join(Account, JournalEntryLine.account_id == Account.id)
        .filter(JournalEntryLine.journal_entry_id.in_(je_ids))
        .all()
    )
    for line, acct in lines:
        net = Decimal(str(line.debit)) - Decimal(str(line.credit))
        t = acct.account_type
        if t in _INCOME_TYPES:
            ni -= net
        elif t in _EXPENSE_TYPES:
            ni += net
        if t in _ASSET_TYPES:
            asset += net
        elif t in _LIABILITY_TYPES:
            liability += net
        elif t in _EQUITY_TYPES:
            equity += net
    return AdjustmentImpact(
        ni_impact=ni,
        ebitda_impact=ni,
        asset_impact=asset,
        liability_impact=liability,
        equity_impact=equity,
    )


@router.get("/advisor-scenarios", response_model=list[AdvisorScenarioOut])
def list_advisor_scenarios(
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org = _org_from_user(user)
    rows = (
        db.query(AdvisorScenario)
        .filter(AdvisorScenario.organization_id == org)
        .order_by(AdvisorScenario.created_at.desc())
        .all()
    )
    return [_scenario_to_out(db, s) for s in rows]


@router.post("/advisor-scenarios", response_model=AdvisorScenarioOut, status_code=201)
def create_advisor_scenario(
    body: AdvisorScenarioCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    if body.scenario_type not in ADVISOR_SCENARIO_TYPES:
        raise HTTPException(status_code=422, detail="Invalid scenario_type")
    org = _org_from_user(user)
    scen = AdvisorScenario(
        organization_id=org,
        name=body.name,
        scenario_type=body.scenario_type,
        description=body.description,
        created_by_user_id=getattr(user, "id", None),
    )
    db.add(scen)
    db.flush()
    db.refresh(scen)
    return _scenario_to_out(db, scen)


# NOTE: /compare must be defined before /{scenario_id} to avoid path collision
@router.get("/advisor-scenarios/compare", response_model=ScenarioComparisonResult)
def compare_advisor_scenarios(
    scenario_ids: list[int] = Query(...),
    entity_id: int = Query(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    results: list[ScenarioImpactResult] = []
    for scen_id in scenario_ids:
        scen = db.get(AdvisorScenario, scen_id)
        if scen is None:
            continue
        included_pkg_ids = [asp.package_id for asp in scen.packages if asp.included]
        je_ids: list[int] = []
        for pkg_id in included_pkg_ids:
            members = (
                db.query(AdjustmentPackageMembership.journal_entry_id)
                .filter(AdjustmentPackageMembership.package_id == pkg_id)
                .all()
            )
            je_ids.extend(m[0] for m in members)
        if je_ids:
            entity_je_ids = [
                row[0] for row in (
                    db.query(JournalEntry.id)
                    .filter(JournalEntry.id.in_(je_ids), JournalEntry.entity_id == entity_id)
                    .all()
                )
            ]
        else:
            entity_je_ids = []
        impact = _compute_impact_for_jes(db, entity_je_ids)
        pkg_names = [
            db.get(AdjustmentPackage, asp.package_id).name
            for asp in scen.packages
            if asp.included and db.get(AdjustmentPackage, asp.package_id)
        ]
        results.append(ScenarioImpactResult(
            scenario_id=scen.id,
            scenario_name=scen.name,
            packages=pkg_names,
            impact=impact,
        ))
    return ScenarioComparisonResult(scenarios=results)


@router.get("/advisor-scenarios/{scenario_id}", response_model=AdvisorScenarioOut)
def get_advisor_scenario(
    scenario_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    scen = db.get(AdvisorScenario, scenario_id)
    if scen is None:
        raise HTTPException(status_code=404, detail="Scenario not found")
    return _scenario_to_out(db, scen)


@router.put("/advisor-scenarios/{scenario_id}", response_model=AdvisorScenarioOut)
def update_advisor_scenario(
    scenario_id: int,
    body: AdvisorScenarioUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    scen = db.get(AdvisorScenario, scenario_id)
    if scen is None:
        raise HTTPException(status_code=404, detail="Scenario not found")
    if body.name is not None:
        scen.name = body.name
    if body.scenario_type is not None:
        if body.scenario_type not in ADVISOR_SCENARIO_TYPES:
            raise HTTPException(status_code=422, detail="Invalid scenario_type")
        scen.scenario_type = body.scenario_type
    if body.description is not None:
        scen.description = body.description
    scen.updated_at = datetime.datetime.utcnow()
    db.flush()
    db.refresh(scen)
    return _scenario_to_out(db, scen)


@router.delete("/advisor-scenarios/{scenario_id}", status_code=204)
def delete_advisor_scenario(
    scenario_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    scen = db.get(AdvisorScenario, scenario_id)
    if scen is None:
        raise HTTPException(status_code=404, detail="Scenario not found")
    db.delete(scen)


@router.post("/advisor-scenarios/{scenario_id}/packages", response_model=dict, status_code=201)
def add_package_to_scenario(
    scenario_id: int,
    body: AdvisorScenarioPackageIn,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    scen = db.get(AdvisorScenario, scenario_id)
    if scen is None:
        raise HTTPException(status_code=404, detail="Scenario not found")
    if db.get(AdjustmentPackage, body.package_id) is None:
        raise HTTPException(status_code=404, detail="Package not found")
    existing = (
        db.query(AdvisorScenarioPackage)
        .filter_by(scenario_id=scenario_id, package_id=body.package_id)
        .first()
    )
    if existing:
        existing.included = body.included
        existing.include_order = body.include_order
        db.flush()
        return {"scenario_id": scenario_id, "package_id": body.package_id, "updated": True}
    asp = AdvisorScenarioPackage(
        scenario_id=scenario_id,
        package_id=body.package_id,
        included=body.included,
        include_order=body.include_order,
    )
    db.add(asp)
    db.flush()
    return {"scenario_id": scenario_id, "package_id": body.package_id, "added": True}


@router.delete("/advisor-scenarios/{scenario_id}/packages/{package_id}", status_code=204)
def remove_package_from_scenario(
    scenario_id: int,
    package_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    asp = (
        db.query(AdvisorScenarioPackage)
        .filter_by(scenario_id=scenario_id, package_id=package_id)
        .first()
    )
    if asp is None:
        raise HTTPException(status_code=404, detail="Package not in scenario")
    db.delete(asp)


@router.patch(
    "/advisor-scenarios/{scenario_id}/packages/{package_id}/toggle",
    response_model=AdvisorScenarioOut,
)
def toggle_package_in_scenario(
    scenario_id: int,
    package_id: int,
    body: PackageToggleIn,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    asp = (
        db.query(AdvisorScenarioPackage)
        .filter_by(scenario_id=scenario_id, package_id=package_id)
        .first()
    )
    if asp is None:
        raise HTTPException(status_code=404, detail="Package not in scenario")
    asp.included = body.included
    db.flush()
    scen = db.get(AdvisorScenario, scenario_id)
    db.refresh(scen)
    return _scenario_to_out(db, scen)


@router.get("/advisor-scenarios/{scenario_id}/export")
def export_advisor_scenario(
    scenario_id: int,
    entity_id: int = Query(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    scen = db.get(AdvisorScenario, scenario_id)
    if scen is None:
        raise HTTPException(status_code=404, detail="Scenario not found")

    included_pkg_ids = [asp.package_id for asp in scen.packages if asp.included]
    je_ids: list[int] = []
    for pkg_id in included_pkg_ids:
        members = (
            db.query(AdjustmentPackageMembership.journal_entry_id)
            .filter(AdjustmentPackageMembership.package_id == pkg_id)
            .all()
        )
        je_ids.extend(m[0] for m in members)

    entity_je_ids: list[int] = []
    if je_ids:
        entity_je_ids = [
            row[0] for row in (
                db.query(JournalEntry.id)
                .filter(JournalEntry.id.in_(je_ids), JournalEntry.entity_id == entity_id)
                .all()
            )
        ]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Adjustments"

    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    WHITE_FONT = Font(bold=True, color="FFFFFF")

    headers = ["JE Number", "Date", "Description", "Amount (Dr)",
               "NI Impact", "Asset Impact", "Liability Impact", "Equity Impact",
               "Materiality", "Status"]
    col_widths = [16, 12, 50, 14, 14, 14, 16, 14, 14, 10]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(1, ci, h)
        c.font = WHITE_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[c.column_letter].width = w

    for ri, je_id in enumerate(entity_je_ids, 2):
        je = db.get(JournalEntry, je_id)
        if je is None:
            continue
        impact = _compute_impact(db, je_id)
        amount = _total_debit(db, je_id)
        ws.cell(ri, 1, je.je_number)
        ws.cell(ri, 2, str(je.entry_date))
        ws.cell(ri, 3, je.description)
        ws.cell(ri, 4, float(amount))
        ws.cell(ri, 5, float(impact.ni_impact))
        ws.cell(ri, 6, float(impact.asset_impact))
        ws.cell(ri, 7, float(impact.liability_impact))
        ws.cell(ri, 8, float(impact.equity_impact))
        ws.cell(ri, 9, je.materiality or "")
        ws.cell(ri, 10, je.status)

    buf = BytesIO()
    wb.save(buf)
    slug = scen.name.replace(" ", "_").lower()
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="scenario_{slug}.xlsx"'},
    )
