"""
Deliverable Workspace API — unified package management for engagement outputs.

Provides: package CRUD, package items, advisor memos, dashboard metrics,
and export stubs (Excel).
"""
from __future__ import annotations

import datetime
import io
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.api.deps import get_db, get_current_user
from app.api.schemas import (
    DeliverableDashboard,
    DeliverableMemoCreate,
    DeliverableMemoOut,
    DeliverableMemoUpdate,
    DeliverablePackageCreate,
    DeliverablePackageItemCreate,
    DeliverablePackageItemOut,
    DeliverablePackageOut,
    DeliverablePackageUpdate,
)
from app.models.deliverable_workspace import (
    DeliverableMemo,
    DeliverablePackage,
    DeliverablePackageItem,
)

router = APIRouter(prefix="/deliverable-workspace", tags=["deliverable-workspace"])

_VALID_TYPES = {"audit", "advisor", "management", "tax", "qoe", "close", "lender", "custom"}
_VALID_STATUSES = {"draft", "internal_review", "client_review", "finalized", "archived"}
_VALID_ITEM_TYPES = {
    "journal_entry", "adjustment_set", "report",
    "financial_statement", "document", "reconciliation", "workpaper",
}


def _org_from_user(user) -> str:
    return getattr(user, "organization_id", None) or "default-org"


def _pkg_or_404(db: Session, pkg_id: int, org_id: str) -> DeliverablePackage:
    pkg = db.query(DeliverablePackage).filter(
        DeliverablePackage.id == pkg_id,
        DeliverablePackage.organization_id == org_id,
    ).first()
    if not pkg:
        raise HTTPException(status_code=404, detail="Package not found")
    return pkg


def _pkg_out(db: Session, pkg: DeliverablePackage) -> DeliverablePackageOut:
    item_count = db.query(func.count(DeliverablePackageItem.id)).filter(
        DeliverablePackageItem.package_id == pkg.id,
    ).scalar() or 0
    memo_count = db.query(func.count(DeliverableMemo.id)).filter(
        DeliverableMemo.package_id == pkg.id,
    ).scalar() or 0
    out = DeliverablePackageOut.model_validate(pkg)
    out.item_count = item_count
    out.memo_count = memo_count
    return out


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@router.get("/dashboard", response_model=DeliverableDashboard)
def get_dashboard(
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    rows = (
        db.query(DeliverablePackage.status, func.count(DeliverablePackage.id))
        .filter(DeliverablePackage.organization_id == org_id)
        .group_by(DeliverablePackage.status)
        .all()
    )
    counts = {r[0]: r[1] for r in rows}
    total = sum(counts.values())
    return DeliverableDashboard(
        total_packages=total,
        draft_count=counts.get("draft", 0),
        internal_review_count=counts.get("internal_review", 0),
        client_review_count=counts.get("client_review", 0),
        finalized_count=counts.get("finalized", 0),
        archived_count=counts.get("archived", 0),
    )


# ---------------------------------------------------------------------------
# Packages
# ---------------------------------------------------------------------------

@router.get("/packages", response_model=list[DeliverablePackageOut])
def list_packages(
    status: Optional[str] = Query(None),
    package_type: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    q = db.query(DeliverablePackage).filter(DeliverablePackage.organization_id == org_id)
    if status:
        q = q.filter(DeliverablePackage.status == status)
    if package_type:
        q = q.filter(DeliverablePackage.package_type == package_type)
    pkgs = q.order_by(DeliverablePackage.created_at.desc()).all()
    return [_pkg_out(db, p) for p in pkgs]


@router.post("/packages", response_model=DeliverablePackageOut, status_code=201)
def create_package(
    body: DeliverablePackageCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    if body.package_type not in _VALID_TYPES:
        raise HTTPException(status_code=422, detail=f"Invalid package_type: {body.package_type}")
    org_id = _org_from_user(user)
    pkg = DeliverablePackage(
        organization_id=org_id,
        name=body.name,
        package_type=body.package_type,
        status="draft",
        description=body.description,
        owner=body.owner,
        created_by_user_id=getattr(user, "id", None),
    )
    db.add(pkg)
    db.commit()
    db.refresh(pkg)
    return _pkg_out(db, pkg)


@router.put("/packages/{pkg_id}", response_model=DeliverablePackageOut)
def update_package(
    pkg_id: int,
    body: DeliverablePackageUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    pkg = _pkg_or_404(db, pkg_id, org_id)
    if body.name is not None:
        pkg.name = body.name
    if body.package_type is not None:
        if body.package_type not in _VALID_TYPES:
            raise HTTPException(status_code=422, detail=f"Invalid package_type: {body.package_type}")
        pkg.package_type = body.package_type
    if body.status is not None:
        if body.status not in _VALID_STATUSES:
            raise HTTPException(status_code=422, detail=f"Invalid status: {body.status}")
        pkg.status = body.status
    if body.description is not None:
        pkg.description = body.description
    if body.owner is not None:
        pkg.owner = body.owner
    pkg.updated_at = datetime.datetime.utcnow()
    db.commit()
    db.refresh(pkg)
    return _pkg_out(db, pkg)


@router.delete("/packages/{pkg_id}", status_code=204)
def delete_package(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    pkg = _pkg_or_404(db, pkg_id, org_id)
    db.delete(pkg)
    db.commit()


@router.post("/packages/{pkg_id}/clone", response_model=DeliverablePackageOut, status_code=201)
def clone_package(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    src = _pkg_or_404(db, pkg_id, org_id)
    clone = DeliverablePackage(
        organization_id=org_id,
        name=f"{src.name} (Copy)",
        package_type=src.package_type,
        status="draft",
        description=src.description,
        owner=src.owner,
        created_by_user_id=getattr(user, "id", None),
    )
    db.add(clone)
    db.flush()
    for item in db.query(DeliverablePackageItem).filter(DeliverablePackageItem.package_id == src.id).all():
        db.add(DeliverablePackageItem(
            package_id=clone.id,
            item_type=item.item_type,
            item_ref=item.item_ref,
            item_label=item.item_label,
        ))
    db.commit()
    db.refresh(clone)
    return _pkg_out(db, clone)


# ---------------------------------------------------------------------------
# Package items
# ---------------------------------------------------------------------------

@router.get("/packages/{pkg_id}/items", response_model=list[DeliverablePackageItemOut])
def list_items(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    return db.query(DeliverablePackageItem).filter(
        DeliverablePackageItem.package_id == pkg_id,
    ).order_by(DeliverablePackageItem.added_at).all()


@router.post("/packages/{pkg_id}/items", response_model=DeliverablePackageItemOut, status_code=201)
def add_item(
    pkg_id: int,
    body: DeliverablePackageItemCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    if body.item_type not in _VALID_ITEM_TYPES:
        raise HTTPException(status_code=422, detail=f"Invalid item_type: {body.item_type}")
    item = DeliverablePackageItem(
        package_id=pkg_id,
        item_type=body.item_type,
        item_ref=body.item_ref,
        item_label=body.item_label,
        added_by_user_id=getattr(user, "id", None),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.delete("/packages/{pkg_id}/items/{item_id}", status_code=204)
def remove_item(
    pkg_id: int,
    item_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    item = db.query(DeliverablePackageItem).filter(
        DeliverablePackageItem.id == item_id,
        DeliverablePackageItem.package_id == pkg_id,
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    db.delete(item)
    db.commit()


# ---------------------------------------------------------------------------
# Memos
# ---------------------------------------------------------------------------

@router.get("/packages/{pkg_id}/memos", response_model=list[DeliverableMemoOut])
def list_memos(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    return db.query(DeliverableMemo).filter(
        DeliverableMemo.package_id == pkg_id,
    ).order_by(DeliverableMemo.created_at).all()


@router.post("/packages/{pkg_id}/memos", response_model=DeliverableMemoOut, status_code=201)
def create_memo(
    pkg_id: int,
    body: DeliverableMemoCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    memo = DeliverableMemo(
        package_id=pkg_id,
        issue=body.issue,
        observation=body.observation,
        recommendation=body.recommendation,
        client_response=body.client_response,
        status=body.status,
        created_by_user_id=getattr(user, "id", None),
    )
    db.add(memo)
    db.commit()
    db.refresh(memo)
    return memo


@router.put("/packages/{pkg_id}/memos/{memo_id}", response_model=DeliverableMemoOut)
def update_memo(
    pkg_id: int,
    memo_id: int,
    body: DeliverableMemoUpdate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    memo = db.query(DeliverableMemo).filter(
        DeliverableMemo.id == memo_id,
        DeliverableMemo.package_id == pkg_id,
    ).first()
    if not memo:
        raise HTTPException(status_code=404, detail="Memo not found")
    for field in ("issue", "observation", "recommendation", "client_response", "status"):
        val = getattr(body, field)
        if val is not None:
            setattr(memo, field, val)
    memo.updated_at = datetime.datetime.utcnow()
    db.commit()
    db.refresh(memo)
    return memo


@router.delete("/packages/{pkg_id}/memos/{memo_id}", status_code=204)
def delete_memo(
    pkg_id: int,
    memo_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    _pkg_or_404(db, pkg_id, org_id)
    memo = db.query(DeliverableMemo).filter(
        DeliverableMemo.id == memo_id,
        DeliverableMemo.package_id == pkg_id,
    ).first()
    if not memo:
        raise HTTPException(status_code=404, detail="Memo not found")
    db.delete(memo)
    db.commit()


# ---------------------------------------------------------------------------
# Export stubs — generate Excel workbooks
# ---------------------------------------------------------------------------

def _wb_to_stream(wb: openpyxl.Workbook) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/packages/{pkg_id}/export/excel")
def export_package_excel(
    pkg_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    org_id = _org_from_user(user)
    pkg = _pkg_or_404(db, pkg_id, org_id)
    items = db.query(DeliverablePackageItem).filter(
        DeliverablePackageItem.package_id == pkg_id,
    ).all()
    memos = db.query(DeliverableMemo).filter(
        DeliverableMemo.package_id == pkg_id,
    ).all()

    wb = openpyxl.Workbook()

    # Cover sheet
    ws_cover = wb.active
    ws_cover.title = "Package"
    ws_cover["A1"] = pkg.name
    ws_cover["A2"] = f"Type: {pkg.package_type}"
    ws_cover["A3"] = f"Status: {pkg.status}"
    ws_cover["A4"] = f"Owner: {pkg.owner or '—'}"
    ws_cover["A5"] = f"Generated: {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"

    # Contents sheet
    ws_items = wb.create_sheet("Contents")
    ws_items.append(["Item Type", "Reference", "Label", "Added"])
    for item in items:
        ws_items.append([
            item.item_type,
            item.item_ref,
            item.item_label or "",
            item.added_at.strftime("%Y-%m-%d") if item.added_at else "",
        ])

    # Memos sheet
    ws_memos = wb.create_sheet("Advisor Memos")
    ws_memos.append(["Issue", "Observation", "Recommendation", "Client Response", "Status"])
    for memo in memos:
        ws_memos.append([
            memo.issue or "",
            memo.observation or "",
            memo.recommendation or "",
            memo.client_response or "",
            memo.status,
        ])

    buf = _wb_to_stream(wb)
    safe_name = pkg.name.replace(" ", "_")[:50]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}.xlsx"'},
    )


@router.get("/exports/adjustment-listing/excel")
def export_adjustment_listing(
    entity_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    from app.models.journal_entry import JournalEntry
    from app.models.journal_entry_line import JournalEntryLine
    from sqlalchemy import func as sqlfunc

    q = db.query(JournalEntry).filter(JournalEntry.status.in_(["draft", "posted"]))
    if entity_id:
        q = q.filter(JournalEntry.entity_id == entity_id)
    entries = q.order_by(JournalEntry.entry_date).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjustment Listing"
    ws.append(["JE Number", "Date", "Entity", "Description", "Source", "Status", "Overlay Group", "Materiality", "Total Debit"])
    for je in entries:
        total_dr = db.query(sqlfunc.sum(JournalEntryLine.debit)).filter(
            JournalEntryLine.journal_entry_id == je.id,
        ).scalar() or 0
        ws.append([
            je.je_number, str(je.entry_date), je.entity_id,
            je.description, je.source, je.status,
            je.overlay_group or "", je.materiality or "", float(total_dr),
        ])

    buf = _wb_to_stream(wb)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="adjustment_listing.xlsx"'},
    )
