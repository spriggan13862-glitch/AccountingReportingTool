"""
M15 proof-point tests: report-run model, Excel export, and package generation.

Tests
-----
1.  Report runs persist correctly (create → status=pending, correct fields)
2.  Excel files generate correctly (execute → non-empty bytes, .xlsx)
3.  Workbook tabs generate correctly (verify expected sheet names)
4.  Storage paths generate correctly (follows org_slug/reports/ convention)
5.  Report exports attach correctly (Document record created on completion)
6.  Organization isolation enforced (list_report_runs scoped to org)
7.  Workflow summaries captured in workbook (sheet present when tasks exist)
8.  Issue summaries captured in workbook (Issues sheet present when issues exist)
9.  Reruns create new run from same parameters and complete successfully
10. Deterministic exports produce stable checksums (same data → same bytes)
"""

import datetime
import io
import zipfile
from decimal import Decimal

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import openpyxl

from app.database import Base
from app.models.entity import Entity
from app.models.scenario import Scenario
from app.models.account import Account
from app.schemas.journal_entry import JournalEntryCreate, JournalEntryLineCreate
from app.services.accounting_period_service import create_period
from app.services.journal_entry_service import post_journal_entry
from app.services.organization_service import create_organization, seed_default_roles
from app.services.user_service import assign_role, create_user
from app.services.report_service import (
    ReportRunStateError,
    ReportValidationError,
    create_report_run,
    execute_report_run,
    get_report_run_or_raise,
    list_report_runs,
    rerun_report,
)
from app.services.export_service import (
    WorkbookMeta,
    build_trial_balance_workbook,
    build_fs_workbook,
    workbook_to_bytes,
)
from app.services.workflow_service import create_task, create_issue


# ---------------------------------------------------------------------------
# In-memory storage for tests
# ---------------------------------------------------------------------------

class MemoryStorage:
    def __init__(self):
        self._store: dict[str, bytes] = {}

    def save(self, path: str, data: bytes) -> None:
        self._store[path] = data

    def load(self, path: str) -> bytes:
        return self._store[path]

    def delete(self, path: str) -> None:
        self._store.pop(path, None)

    def exists(self, path: str) -> bool:
        return path in self._store


# ---------------------------------------------------------------------------
# Module-scoped fixture
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def seeded_session():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    db = Session()

    org_a = create_organization(db, name="Org Alpha", slug="alpha")
    org_b = create_organization(db, name="Org Beta", slug="beta")
    seed_default_roles(db)

    admin = create_user(db, org_a.id, "admin@rpt.com", "Admin")
    assign_role(db, admin.id, "admin", org_a.id)

    accountant = create_user(db, org_a.id, "acct@rpt.com", "Accountant")
    assign_role(db, accountant.id, "accountant", org_a.id)

    entity_a = Entity(
        code="RA", name="Report Entity A",
        entity_type="operating", currency="USD",
        organization_id=org_a.id,
    )
    db.add(entity_a)
    db.flush()

    cash = Account(
        entity_id=entity_a.id, account_number="1000",
        account_name="Cash", account_type="asset", normal_balance="debit",
    )
    rev = Account(
        entity_id=entity_a.id, account_number="4000",
        account_name="Revenue", account_type="revenue", normal_balance="credit",
    )
    re_acc = Account(
        entity_id=entity_a.id, account_number="3900",
        account_name="RE", account_type="equity", normal_balance="credit",
    )
    for acc in (cash, rev, re_acc):
        db.add(acc)
    db.flush()

    scenario = Scenario(code="ACT", name="Actual", scenario_type="actual")
    db.add(scenario)
    db.flush()

    je = post_journal_entry(
        db,
        JournalEntryCreate(
            je_number="JE-RPT-001",
            entry_date=datetime.date(2024, 3, 31),
            entity_id=entity_a.id,
            scenario_id=scenario.id,
            description="March revenue",
            source="test",
            lines=[
                JournalEntryLineCreate(
                    line_number=1, account_id=cash.id, entity_id=entity_a.id,
                    debit=Decimal("5000"), credit=Decimal("0"),
                ),
                JournalEntryLineCreate(
                    line_number=2, account_id=rev.id, entity_id=entity_a.id,
                    debit=Decimal("0"), credit=Decimal("5000"),
                ),
            ],
        ),
        acting_user=accountant,
    )

    period = create_period(
        db, entity_id=entity_a.id,
        period_name="March 2024",
        start_date=datetime.date(2024, 3, 1),
        end_date=datetime.date(2024, 3, 31),
        fiscal_year=2024, fiscal_period=3,
    )

    db.commit()

    storage = MemoryStorage()

    yield db, storage, {
        "org_a": org_a, "org_b": org_b,
        "admin": admin, "accountant": accountant,
        "entity_a": entity_a,
        "cash": cash, "rev": rev, "re_acc": re_acc,
        "scenario": scenario,
        "je": je,
        "period": period,
    }

    db.close()
    Base.metadata.drop_all(engine)


@pytest.fixture
def s(seeded_session):
    db, storage, d = seeded_session
    db.begin_nested()
    yield db, storage, d
    db.rollback()


# ---------------------------------------------------------------------------
# Test 1: Report runs persist correctly
# ---------------------------------------------------------------------------

def test_report_run_persists(s):
    db, storage, d = s
    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="trial_balance",
        output_format="xlsx",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
        acting_user=d["accountant"],
    )
    assert run.id is not None
    assert run.status == "pending"
    assert run.report_type == "trial_balance"
    assert run.output_format == "xlsx"
    assert run.entity_id == d["entity_a"].id
    assert run.created_by_user_id == d["accountant"].id
    assert run.storage_path is None
    assert run.completed_at is None


def test_report_run_invalid_type(s):
    db, storage, d = s
    with pytest.raises(ReportValidationError, match="Unknown report_type"):
        create_report_run(
            db, organization_id=d["org_a"].id,
            report_type="nonexistent_report",
        )


# ---------------------------------------------------------------------------
# Test 2: Excel files generate correctly
# ---------------------------------------------------------------------------

def test_excel_generates_non_empty_bytes(s):
    db, storage, d = s
    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="trial_balance",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
        acting_user=d["accountant"],
    )
    db.flush()
    run = execute_report_run(db, run.id, storage, "alpha")
    assert run.status == "completed"
    assert run.storage_path is not None
    file_bytes = storage.load(run.storage_path)
    assert len(file_bytes) > 0
    # Valid zip (xlsx is a zip)
    assert zipfile.is_zipfile(io.BytesIO(file_bytes))


# ---------------------------------------------------------------------------
# Test 3: Workbook tabs generate correctly
# ---------------------------------------------------------------------------

def test_workbook_sheet_names(s):
    db, storage, d = s
    meta = WorkbookMeta(
        report_type="trial_balance",
        entity_name="Test Entity",
        as_of_date=datetime.date(2024, 3, 31),
    )
    from app.services.reporting_service import get_trial_balance
    tb_rows = get_trial_balance(
        db, d["entity_a"].id,
        datetime.date(2024, 3, 31),
        [d["scenario"].id],
    )
    wb = build_trial_balance_workbook(meta, tb_rows)
    sheet_names = wb.sheetnames
    assert "Cover" in sheet_names
    assert "Metadata" in sheet_names
    assert "Trial Balance" in sheet_names


def test_fs_workbook_sheet_names(s):
    db, storage, d = s
    meta = WorkbookMeta(report_type="balance_sheet", entity_name="Test Entity")
    wb = build_fs_workbook(meta, [], [])
    sheet_names = wb.sheetnames
    assert "Cover" in sheet_names
    assert "Metadata" in sheet_names
    assert "Trial Balance" in sheet_names


# ---------------------------------------------------------------------------
# Test 4: Storage paths follow convention
# ---------------------------------------------------------------------------

def test_storage_path_convention(s):
    db, storage, d = s
    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="trial_balance",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
    )
    db.flush()
    run = execute_report_run(db, run.id, storage, "alpha")
    assert run.storage_path.startswith("alpha/reports/")
    assert run.storage_path.endswith(".xlsx")


# ---------------------------------------------------------------------------
# Test 5: Document record created on completion
# ---------------------------------------------------------------------------

def test_document_record_created(s):
    db, storage, d = s
    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="trial_balance",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
    )
    db.flush()
    run = execute_report_run(db, run.id, storage, "alpha")
    assert run.status == "completed"
    assert run.generated_document_id is not None

    from app.models.document import Document
    doc = db.get(Document, run.generated_document_id)
    assert doc is not None
    assert doc.organization_id == d["org_a"].id
    assert doc.document_type == "workpaper"
    assert len(doc.checksum_sha256) == 64


# ---------------------------------------------------------------------------
# Test 6: Organization isolation
# ---------------------------------------------------------------------------

def test_org_isolation(s):
    db, storage, d = s
    run_a = create_report_run(
        db, organization_id=d["org_a"].id, report_type="trial_balance",
    )
    run_b = create_report_run(
        db, organization_id=d["org_b"].id, report_type="trial_balance",
    )
    db.flush()

    runs_a = list_report_runs(db, d["org_a"].id)
    runs_b = list_report_runs(db, d["org_b"].id)

    ids_a = {r.id for r in runs_a}
    ids_b = {r.id for r in runs_b}

    assert run_a.id in ids_a
    assert run_b.id not in ids_a
    assert run_b.id in ids_b
    assert run_a.id not in ids_b


# ---------------------------------------------------------------------------
# Test 7: Workflow summaries appear in workbook
# ---------------------------------------------------------------------------

def test_workflow_summary_in_workbook(s):
    db, storage, d = s
    create_task(
        db, d["org_a"].id,
        task_type="close_task",
        title="Q1 Close Task",
        acting_user=d["admin"],
    )
    db.flush()

    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="close_package",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
    )
    db.flush()
    run = execute_report_run(db, run.id, storage, "alpha")
    assert run.status == "completed"

    import json
    wf = json.loads(run.workflow_summary_json)
    assert "tasks" in wf
    assert len(wf["tasks"]) > 0

    file_bytes = storage.load(run.storage_path)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    assert "Workflow" in wb.sheetnames


# ---------------------------------------------------------------------------
# Test 8: Issue summaries appear in workbook
# ---------------------------------------------------------------------------

def test_issue_summary_in_workbook(s):
    db, storage, d = s
    create_issue(
        db, d["org_a"].id,
        issue_code="TB_OUT_OF_BALANCE",
        severity="error",
        title="TB out of balance for Q1",
        acting_user=d["admin"],
    )
    db.flush()

    run = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="audit_support_package",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
    )
    db.flush()
    run = execute_report_run(db, run.id, storage, "alpha")
    assert run.status == "completed"

    import json
    wf = json.loads(run.workflow_summary_json)
    assert "issues" in wf
    assert len(wf["issues"]) > 0

    file_bytes = storage.load(run.storage_path)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    assert "Issues" in wb.sheetnames


# ---------------------------------------------------------------------------
# Test 9: Reruns work correctly
# ---------------------------------------------------------------------------

def test_rerun_creates_new_run(s):
    db, storage, d = s
    original = create_report_run(
        db,
        organization_id=d["org_a"].id,
        report_type="trial_balance",
        entity_id=d["entity_a"].id,
        scenario_ids=[d["scenario"].id],
        parameters={"as_of_date": "2024-03-31"},
    )
    db.flush()
    original = execute_report_run(db, original.id, storage, "alpha")
    assert original.status == "completed"

    new_run = rerun_report(db, original.id, storage, "alpha")
    assert new_run.id != original.id
    assert new_run.status == "completed"
    assert new_run.report_type == original.report_type
    assert new_run.entity_id == original.entity_id
    assert new_run.storage_path != original.storage_path


# ---------------------------------------------------------------------------
# Test 10: Deterministic exports produce stable checksums
# ---------------------------------------------------------------------------

def test_deterministic_export_stable(s):
    db, storage, d = s
    from app.services.reporting_service import get_trial_balance

    tb_rows = get_trial_balance(
        db, d["entity_a"].id,
        datetime.date(2024, 3, 31),
        [d["scenario"].id],
    )
    fixed_time = datetime.datetime(2024, 1, 1, 0, 0, 0)
    meta = WorkbookMeta(
        report_type="trial_balance",
        entity_name="Stable Entity",
        as_of_date=datetime.date(2024, 3, 31),
        generation_time=fixed_time,
    )

    wb1 = build_trial_balance_workbook(meta, tb_rows)
    bytes1 = workbook_to_bytes(wb1, generation_time=fixed_time, deterministic=True)

    wb2 = build_trial_balance_workbook(meta, tb_rows)
    bytes2 = workbook_to_bytes(wb2, generation_time=fixed_time, deterministic=True)

    assert bytes1 == bytes2
    assert len(bytes1) > 0
